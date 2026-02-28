import os
import re
import json
import time
import getpass
import shutil
import unicodedata
from dataclasses import dataclass, asdict
from datetime import datetime, timedelta
from pathlib import Path
from difflib import get_close_matches
from typing import Optional, Dict, List, Tuple, Any

import openpyxl
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill, Font

from prompt_toolkit import prompt
from prompt_toolkit.completion import Completer, Completion
from prompt_toolkit.application import Application, get_app
from prompt_toolkit.key_binding import KeyBindings
from prompt_toolkit.layout import Layout
from prompt_toolkit.layout.containers import HSplit, VSplit, Window
from prompt_toolkit.layout.controls import FormattedTextControl
from prompt_toolkit.styles import Style

from rich.console import Console
from rich.table import Table
from rich.panel import Panel
from rich.columns import Columns
from rich.align import Align
from rich import box

import keyboard
import win32com.client

from selenium import webdriver
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.chrome.service import Service
from selenium.common.exceptions import (
    TimeoutException,
    WebDriverException,
    StaleElementReferenceException,
    ElementClickInterceptedException,
)
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.support.ui import WebDriverWait

console = Console()

FRAME_CACHE: Dict[str, Optional[int]] = {"default": None, "diario": None, "frequencia": None, "quadro": None}

# ==========================================================
# CONFIG
# ==========================================================
ARQUIVO_HORARIOS = "horarios_megatron.json"
ARQUIVO_LEMBRETES = "lembretes_megatron.json"
ARQUIVO_STATUS = "status_turmas_megatron.json"
ARQUIVO_PENDENCIAS = "pendencias_megatron.json"
CACHE_QUADROS = Path("dump_quadros_cache.json")
MEM_PROFS = Path("memoria_professores.json")
MAPEAMENTO_TURMAS = Path("mapeamento_turmas_megatron.json")
ARQUIVO_BANCO_ALUNOS = Path("banco_alunos_megatron.json")
ARQUIVO_PERFIL_SIEPE = Path("perfil_siepe_megatron.json")
ARQUIVO_ROTEIRO = Path("roteiro_usuario_megatron.json")
CRED_FILE = Path("meu_acesso.txt")

SIEPE_URL = "https://siepe.educacao.pe.gov.br/"
SEL_EDUCADORES = (By.ID, "mnSupEdu")
XPATH_LINKS_DETALHE = "//a[contains(@href,'detalheQuadroDeHorario(') or contains(@onclick,'detalheQuadroDeHorario(')]"
ID_CONTAINER_TURMAS = "turmasQuadroDeHorario"
XPATH_ALL_TABELAS_HORARIO = "//table[contains(@class,'ListagemPadrao') and contains(@class,'Horario')]"
DIAS_COL = ["SEG", "TER", "QUA", "QUI", "SEX"]
DIAS_HEADERS_ACCEPT = ["SEG", "TER", "QUA", "QUI", "SEX", "SAB", "DOM"]

HORARIOS_PADRAO = [
    "07:30 - 08:20",
    "08:20 - 09:10",
    "09:10 - 10:00",
    "10:20 - 11:10",
    "11:10 - 12:00",
    "13:00 - 13:50",
    "13:50 - 14:40",
]
DIAS_SEMANA = ["Segunda", "Terça", "Quarta", "Quinta", "Sexta"]

FILL_FALTA = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
FILL_NOTA = PatternFill(start_color="9BC2E6", end_color="9BC2E6", fill_type="solid")

FILL_AVALIACAO = {
    1: PatternFill(start_color="C00000", end_color="C00000", fill_type="solid"),  # vermelho
    2: PatternFill(start_color="F4B183", end_color="F4B183", fill_type="solid"),  # laranja
    3: PatternFill(start_color="C9DA2A", end_color="C9DA2A", fill_type="solid"),  # amarelo lima
    4: PatternFill(start_color="92D050", end_color="92D050", fill_type="solid"),  # verde lima
    5: PatternFill(start_color="385723", end_color="385723", fill_type="solid"),  # verde musgo
}
FONT_BRANCA = Font(color="FFFFFF")
FONT_PRETA = Font(color="000000")

TUI_STYLE = Style.from_dict(
    {
        "title": "bold #ffff00",
        "hint": "#888888",
        "item": "#ffffff",
        "item.selected": "bold #000000 bg:#00ffff",
        "item.marked": "bold #00ff00",
        "border": "#00ffff",
        "search": "#ffaa00",
    }
)

NAV_BACK = "__BACK__"
NAV_MENU = "__MENU__"

lista_turmas = [
    "6A - ENSINO FUNDAMENTAL DE ANOS FINAIS DE 3500H - 6º ANO A",
    "7A - ENSINO FUNDAMENTAL DE ANOS FINAIS DE 3500H - 7º ANO A",
    "8A - ENSINO FUNDAMENTAL DE ANOS FINAIS DE 3500H - 8º ANO A",
    "9A - ENSINO FUNDAMENTAL DE ANOS FINAIS DE 3500H - 9º ANO A",
    "9B - ENSINO FUNDAMENTAL DE ANOS FINAIS DE 3500H - 9º ANO B",
    "AULA ATIVIDADE",
]

banco_alunos = {
    "6A": [
        "ALEFF GABRIEL GOMES DA SILVA",
        "ALISSON TORRES DA SILVA",
        "ALYCIA HADASSA RAINELLE SIQUEIRA DE OLIVEIRA",
        "ARTHUR VINICIUS DE SOUZA MACHADO",
    ],
    "6B": [],
    "7A": [
        "ADRIANA KETHYLLE MELO PERCILIO",
        "ALERRANDRO FERREIRA FIDELIS DA SILVA",
        "ANA JULIA MELO DA SILVA",
        "ARTHUR NUNES BEZERRA",
    ],
    "7B": [],
    "8A": [
        "ANA CLARA DOS SANTOS RAMOS",
        "ANA WYCTORYA DE ALMEIDA SOARES",
        "BRENDA NICOLE DOS SANTOS RAGO",
        "CICERO TINO DA SILVA NETO",
    ],
    "8B": [],
    "9A": [
        "ALEXSANDRO DOS SANTOS DE LIMA",
        "ALLANA VICTORIA SANTOS DE OLIVEIRA",
        "ANA BEATRIZ DA SILVA LEANDRO",
        "ANA CLARA DA SILVA",
    ],
    "9B": [
        "ALEXSANDRO VIDAL DA SILVA",
        "ALISSON ALVES DE LIMA",
        "ANA ROBERTA MOURA MENDES DA SILVA",
        "ATHOS WINICIUS JOSE ALVES DA CUNHA",
    ],
}

apelidos_alunos: Dict[str, Dict[str, str]] = {}

# ==========================================================
# MODELOS
# ==========================================================
@dataclass
class ChamadaRecord:
    turma: str
    data: str
    modo: str
    lista_digitada: List[str]
    faltosos_excel: List[str]
    observacoes: str
    timestamp: str


@dataclass
class PendenciaRecord:
    turma: str
    data: str
    modo: str
    lista_digitada: List[str]
    faltosos_excel: List[str]
    observacoes: str
    timestamp: str
    status: str = "pendente"
    erro: str = ""


@dataclass
class TarefaChecklist:
    texto: str
    concluida: bool = False
    data: str = ""
    horario: str = ""
    prioridade: int = 2


# ==========================================================
# UTIL
# ==========================================================
def log_info(msg: str):
    console.print(f"[cyan]ℹ[/cyan] {msg}")


def log_warn(msg: str):
    console.print(f"[yellow]⚠[/yellow] {msg}")


def log_error(msg: str):
    console.print(f"[red]❌[/red] {msg}")


def normalizar_texto(texto: str) -> str:
    return unicodedata.normalize("NFKD", texto or "").encode("ASCII", "ignore").decode("utf-8").upper().strip()


def agora_slot_info(data_foco: datetime) -> Tuple[Optional[int], Optional[int]]:
    if data_foco.weekday() > 4:
        return None, None
    now = datetime.now()
    hm_now = int(now.strftime("%H%M"))
    idx_atual = None
    idx_prox = None
    for idx, slot in enumerate(HORARIOS_PADRAO):
        m = re.findall(r"(\d{2}):(\d{2})", slot)
        if len(m) != 2:
            continue
        ini = int(m[0][0] + m[0][1])
        fim = int(m[1][0] + m[1][1])
        if ini <= hm_now <= fim:
            idx_atual = idx
        if hm_now < ini and idx_prox is None:
            idx_prox = idx
    if idx_atual is not None and idx_prox is None and idx_atual + 1 < len(HORARIOS_PADRAO):
        idx_prox = idx_atual + 1
    return idx_atual, idx_prox


def fuzzy_sugerir(nome: str, base: List[str], limite=3) -> List[str]:
    n = normalizar_texto(nome)
    base_norm = {normalizar_texto(b): b for b in base}
    candidatos = get_close_matches(n, list(base_norm.keys()), n=limite, cutoff=0.72)
    return [base_norm[c] for c in candidatos]


def fazer_backup_excel(caminho_arquivo: str):
    if os.path.exists(caminho_arquivo):
        try:
            p = Path(caminho_arquivo)
            shutil.copy2(caminho_arquivo, str(p.with_name(f"{p.stem}_BKP{p.suffix}")))
        except Exception:
            pass


def carregar_json(caminho: str, padrao):
    if not os.path.exists(caminho):
        return padrao
    try:
        with open(caminho, "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        return padrao


def salvar_json(caminho: str, dados):
    with open(caminho, "w", encoding="utf-8") as f:
        json.dump(dados, f, ensure_ascii=False, indent=2)


def entrada_cancelada(txt: str) -> bool:
    return normalizar_texto(txt) == "ESC"


def parse_data_pratica(entrada: str, base: datetime) -> datetime:
    s = entrada.strip().upper()
    if s == "H":
        return datetime.now()
    for fmt in ("%d/%m/%Y", "%d/%m"):
        try:
            dt = datetime.strptime(s, fmt)
            return dt.replace(year=base.year) if fmt == "%d/%m" else dt
        except ValueError:
            pass
    if s.isdigit():
        if len(s) == 2:
            return base.replace(day=int(s))
        if len(s) == 4:
            return base.replace(day=int(s[:2]), month=int(s[2:]))
        if len(s) == 8:
            return base.replace(day=int(s[:2]), month=int(s[2:4]), year=int(s[4:]))
    raise ValueError("Formato inválido")


def status_padrao():
    return {"turmas": {}}


def lembretes_padrao():
    return {"turma": {}, "aluno": {}}


def pendencias_padrao():
    return []




def roteiro_padrao():
    return {"preferencias": {}, "trilha": []}


def registrar_trilha_usuario(data_foco: datetime, acao: str, turma: str = "", slot: str = "") -> None:
    roteiro = carregar_json(str(ARQUIVO_ROTEIRO), roteiro_padrao())
    if not isinstance(roteiro, dict):
        roteiro = roteiro_padrao()

    preferencias = roteiro.setdefault("preferencias", {})
    dia = str(data_foco.weekday())
    chave = f"{dia}|{slot}" if slot else dia
    if turma:
        preferencias[chave] = turma

    trilha = roteiro.setdefault("trilha", [])
    trilha.append({
        "ts": datetime.now().isoformat(),
        "acao": acao,
        "dia": dia,
        "slot": slot,
        "turma": turma,
    })
    if len(trilha) > 200:
        trilha[:] = trilha[-200:]

    salvar_json(str(ARQUIVO_ROTEIRO), roteiro)


def obter_turma_prioritaria(data_foco: datetime, grade: dict) -> Tuple[str, str]:
    roteiro = carregar_json(str(ARQUIVO_ROTEIRO), roteiro_padrao())
    preferencias = roteiro.get("preferencias", {}) if isinstance(roteiro, dict) else {}

    agora_h, _prox = agora_slot_info(data_foco)
    slot = agora_h if agora_h in HORARIOS_PADRAO else (HORARIOS_PADRAO[0] if HORARIOS_PADRAO else "")
    dia = str(data_foco.weekday())

    turma = str(preferencias.get(f"{dia}|{slot}", "")).strip().upper()
    if not turma and slot:
        turma = str((grade.get(dia, {}) or {}).get(slot, "")).strip().upper()

    if turma == "AULA ATIVIDADE":
        turma = ""

    return turma, slot


def recarregar_indices_alunos() -> None:
    global BANCO_ALUNOS_NORM
    BANCO_ALUNOS_NORM = {k: [(n, normalizar_texto(n)) for n in v] for k, v in banco_alunos.items()}


def carregar_banco_alunos_completo() -> None:
    """Carrega banco completo de alunos de JSON externo, mantendo fallback interno."""
    global banco_alunos
    if not ARQUIVO_BANCO_ALUNOS.exists():
        salvar_json(str(ARQUIVO_BANCO_ALUNOS), banco_alunos)
        log_warn(f"Arquivo {ARQUIVO_BANCO_ALUNOS.name} criado. Preencha com a lista completa de alunos.")
        recarregar_indices_alunos()
        return

    dados = carregar_json(str(ARQUIVO_BANCO_ALUNOS), {})
    if not isinstance(dados, dict):
        log_warn(f"{ARQUIVO_BANCO_ALUNOS.name} inválido. Mantendo banco interno.")
        recarregar_indices_alunos()
        return

    normalizado = {}
    for turma, nomes in dados.items():
        if not isinstance(nomes, list):
            continue
        itens = []
        seen = set()
        for n in nomes:
            if not isinstance(n, str):
                continue
            nome = n.strip().upper()
            if nome and nome not in seen:
                seen.add(nome)
                itens.append(nome)
        normalizado[str(turma).strip().upper()] = itens

    if normalizado:
        banco_alunos.update(normalizado)
    recarregar_indices_alunos()

def calcular_semana_atual(data_foco: datetime) -> str:
    inicio_aulas = datetime(2026, 2, 3)
    if data_foco < inicio_aulas:
        return "Em Recesso"
    return f"Semana Letiva {((data_foco - inicio_aulas).days // 7) + 1}"


# ==========================================================
# COMPLETERS
# ==========================================================
LISTA_TURMAS_NORM = [(t, normalizar_texto(t)) for t in lista_turmas]
BANCO_ALUNOS_NORM = {k: [(n, normalizar_texto(n)) for n in v] for k, v in banco_alunos.items()}


def _match_priority(nome_norm: str, query_norm: str) -> tuple[int, str]:
    if nome_norm.startswith(query_norm):
        return (0, nome_norm)
    if any(p.startswith(query_norm) for p in nome_norm.split()):
        return (1, nome_norm)
    if query_norm in nome_norm:
        return (2, nome_norm)
    return (9, nome_norm)


class TurmaCompleter(Completer):
    def get_completions(self, document, complete_event):
        q = normalizar_texto(document.text)
        if not q:
            return
        cands = [(t, n) for t, n in LISTA_TURMAS_NORM if _match_priority(n, q)[0] < 9]
        cands.sort(key=lambda x: _match_priority(x[1], q))
        for turma, _ in cands:
            yield Completion(turma, start_position=-len(document.text))


# ==========================================================
# UI BASE
# ==========================================================
def tui_select(title: str, text: str, values: List[Tuple[str, str]], default: Optional[str] = None, allow_back: bool = True, allow_menu: bool = True) -> Optional[str]:
    if not values:
        return None
    idx = next((i for i, (v, _) in enumerate(values) if v == default), 0)
    state = {"idx": idx, "result": None}

    def render_list():
        out = []
        for i, (_v, label) in enumerate(values):
            sty = "class:item.selected" if i == state["idx"] else "class:item"
            pref = " ▶ " if i == state["idx"] else "   "
            out.append((sty, f"{pref}{label}\n"))
        return out

    kb = KeyBindings()

    @kb.add("up")
    def _up(event):
        state["idx"] = (state["idx"] - 1) % len(values)

    @kb.add("down")
    def _down(event):
        state["idx"] = (state["idx"] + 1) % len(values)

    @kb.add("enter")
    def _ent(event):
        state["result"] = values[state["idx"]][0]
        event.app.exit()

    @kb.add("f1")
    def _f1(event):
        if allow_back:
            state["result"] = NAV_BACK
            event.app.exit()

    @kb.add("f12")
    def _f12(event):
        if allow_menu:
            state["result"] = NAV_MENU
            event.app.exit()

    @kb.add("escape")
    def _esc(event):
        state["result"] = NAV_BACK if allow_back else None
        event.app.exit()

    root = HSplit(
        [
            Window(height=1, content=FormattedTextControl([("class:border", "─" * 100)])),
            Window(height=3, content=FormattedTextControl(lambda: [("class:title", f"{title}\n"), ("", f"{text}\n")])),
            Window(height=1, content=FormattedTextControl([("class:border", "─" * 100)])),
            Window(content=FormattedTextControl(render_list)),
            Window(height=1, content=FormattedTextControl([("class:border", "─" * 100)])),
            Window(height=1, content=FormattedTextControl([("class:hint", "↑↓ navegar | ENTER confirmar | F1 voltar | F12 menu | ESC voltar")])),
        ]
    )
    Application(layout=Layout(root), key_bindings=kb, full_screen=True, style=TUI_STYLE).run()
    return state["result"]

def checkbox_selector(title: str, alunos: List[str], modo: str) -> Optional[List[str] | str]:
    filtrado = list(alunos)
    marcados = set()
    idx = 0
    topo = 0
    busca = ""
    state = {"res": None}

    def _altura_lista() -> int:
        try:
            total = get_app().output.get_size().rows
            return max(6, total - 7)
        except Exception:
            return 18

    def _ajustar_janela():
        nonlocal topo
        if not filtrado:
            topo = 0
            return
        max_visiveis = _altura_lista()
        topo = max(0, min(topo, max(0, len(filtrado) - max_visiveis)))
        if idx < topo:
            topo = idx
        if idx >= topo + max_visiveis:
            topo = idx - max_visiveis + 1

    def render_rows():
        rows = []
        rows.append(("class:title", f"{title}\n"))
        rows.append(("class:search", f"Busca: {busca}\n"))
        rows.append(("class:hint", "TAB marca | F2 marcar todos | F3 limpar | PgUp/PgDn rolagem | Enter confirma | F1 voltar | F12 menu | ESC voltar\n"))

        if not filtrado:
            rows.append(("class:hint", "Sem resultados para o filtro.\n"))
            return rows

        _ajustar_janela()
        max_visiveis = _altura_lista()
        fim = min(len(filtrado), topo + max_visiveis)
        faixa = filtrado[topo:fim]

        progresso = 1.0 if len(filtrado) <= max_visiveis else (fim / max(1, len(filtrado)))
        blocos = 12
        preenchido = max(1, min(blocos, int(round(progresso * blocos))))
        barra = "█" * preenchido + "░" * (blocos - preenchido)
        rows.append(("class:hint", f"Selecionados: {len(marcados)} | Exibindo {topo + 1}-{fim}/{len(filtrado)} | Rolagem [{barra}]\n\n"))

        for j, nome in enumerate(faixa, start=topo):
            mk = "✅" if nome in marcados else "☐"
            sty = "class:item.selected" if j == idx else ("class:item.marked" if nome in marcados else "class:item")
            rows.append((sty, f" {mk} {nome}\n"))
        return rows

    def aplicar_filtro():
        nonlocal filtrado, idx, topo
        b = normalizar_texto(busca)
        if not b:
            filtrado = list(alunos)
        else:
            ranqueados = []
            for a in alunos:
                nome_norm = normalizar_texto(a)
                apelido_norm = normalizar_texto(apelidos_alunos.get(a, {}).get("apelido", ""))
                prioridade_nome = _match_priority(nome_norm, b)[0]
                prioridade_apelido = _match_priority(apelido_norm, b)[0] if apelido_norm else 9
                prioridade = min(prioridade_nome, prioridade_apelido)
                if prioridade < 9:
                    ranqueados.append((prioridade, nome_norm, a))

            # prioridade linear: começo do nome > começo de palavra > contains
            ranqueados.sort(key=lambda x: (x[0], x[1]))
            filtrado = [a for _p, _n, a in ranqueados]

            if not filtrado:
                filtrado = fuzzy_sugerir(busca, alunos, limite=5)

        # mantém itens selecionados no final da tabela, preservando a ordem relativa
        nao_marcados = [a for a in filtrado if a not in marcados]
        ja_marcados = [a for a in filtrado if a in marcados]
        filtrado = nao_marcados + ja_marcados

        idx = 0 if filtrado else -1
        topo = 0

    def _reset_busca():
        nonlocal busca
        busca = ""
        aplicar_filtro()

    aplicar_filtro()
    kb = KeyBindings()

    @kb.add("up")
    def _up(event):
        nonlocal idx
        if filtrado:
            idx = (idx - 1) % len(filtrado)
            _ajustar_janela()

    @kb.add("down")
    def _down(event):
        nonlocal idx
        if filtrado:
            idx = (idx + 1) % len(filtrado)
            _ajustar_janela()

    @kb.add("pageup")
    def _pgup(event):
        nonlocal idx
        if filtrado:
            idx = max(0, idx - _altura_lista())
            _ajustar_janela()

    @kb.add("pagedown")
    def _pgdn(event):
        nonlocal idx
        if filtrado:
            idx = min(len(filtrado) - 1, idx + _altura_lista())
            _ajustar_janela()

    @kb.add("home")
    def _home(event):
        nonlocal idx
        if filtrado:
            idx = 0
            _ajustar_janela()

    @kb.add("end")
    def _end(event):
        nonlocal idx
        if filtrado:
            idx = len(filtrado) - 1
            _ajustar_janela()

    @kb.add("tab")
    def _space(event):
        nonlocal marcados, idx
        if not filtrado or idx < 0:
            return

        idx_anterior = idx
        alvo = filtrado[idx]
        if modo == "U":
            marcados = {alvo}
            # mantendo comportamento atual quando há busca ativa
            if busca:
                _reset_busca()
            else:
                aplicar_filtro()
                if filtrado:
                    idx = min(idx_anterior, len(filtrado) - 1)
                    _ajustar_janela()
            return

        if alvo in marcados:
            marcados.remove(alvo)
        else:
            marcados.add(alvo)

        # com busca ativa, mantém comportamento atual (limpa busca e volta para topo)
        if busca:
            _reset_busca()
            return

        # sem busca ativa, segue para o próximo da fila sem voltar ao primeiro item
        aplicar_filtro()
        if filtrado:
            idx = min(idx_anterior, len(filtrado) - 1)
            _ajustar_janela()

    @kb.add("f2")
    def _todos(event):
        nonlocal marcados
        if modo == "U":
            return
        marcados.update(filtrado)
        _reset_busca()

    @kb.add("f3")
    def _limpar(event):
        nonlocal marcados
        if modo == "U":
            marcados.clear()
            _reset_busca()
            return
        marcados.difference_update(filtrado)
        _reset_busca()

    @kb.add("f1")
    def _f1(event):
        state["res"] = NAV_BACK
        event.app.exit()

    @kb.add("f12")
    def _f12(event):
        state["res"] = NAV_MENU
        event.app.exit()

    @kb.add("enter")
    def _enter(event):
        if modo == "U" and len(marcados) != 1:
            return
        state["res"] = sorted(marcados)
        event.app.exit()

    @kb.add("backspace")
    def _bs(event):
        nonlocal busca
        if busca:
            busca = busca[:-1]
            aplicar_filtro()

    @kb.add("escape")
    def _esc(event):
        state["res"] = NAV_BACK
        event.app.exit()

    @kb.add("<any>")
    def _any(event):
        nonlocal busca
        t = event.data
        if not t or not t.isprintable() or t in {"\n", "\r", "\t"}:
            return
        busca += t
        aplicar_filtro()

    root = HSplit([Window(content=FormattedTextControl(render_rows))])
    Application(layout=Layout(root), key_bindings=kb, full_screen=True, style=TUI_STYLE).run()
    return state["res"]

def modal_erro_siepe(msg: str) -> str:
    return tui_select(
        "❌ Falha na automação SIEPE",
        msg,
        [("retry", "Tentar novamente"), ("pend", "Salvar pendência"), ("cancel", "Cancelar")],
        default="retry",
    ) or "cancel"


# ==========================================================
# STORAGE
# ==========================================================
def registrar_status_chamada(status: dict, turma: str, data_str: str, slot: str, presentes: int, faltosos: int, media_turma: Optional[float], avaliacao_nivel: int = 0, avaliacao_label: str = ""):
    t = status.setdefault("turmas", {}).setdefault(turma, {}).setdefault(data_str, {})
    t[slot] = {
        "presentes": presentes,
        "faltosos": faltosos,
        "chamada_feita": True,
        "media_turma": media_turma,
        "avaliacao_nivel": int(avaliacao_nivel or 0),
        "avaliacao_label": str(avaliacao_label or "").strip(),
    }


def slot_chamada_feita(status: dict, turma: str, data_str: str, slot: str) -> bool:
    return bool(status.get("turmas", {}).get(turma, {}).get(data_str, {}).get(slot, {}).get("chamada_feita"))

def _estilo_avaliacao_texto(nivel: int, texto: str) -> str:
    estilos = {
        1: "white on red",
        2: "black on #f4b183",
        3: "black on #c9da2a",
        4: "black on #92d050",
        5: "white on #385723",
    }
    st = estilos.get(int(nivel or 0), "")
    return f"[{st}]{texto}[/]" if st else texto



def adicionar_pendencia(reg: PendenciaRecord):
    pend = carregar_json(ARQUIVO_PENDENCIAS, pendencias_padrao())
    pend.append(asdict(reg))
    salvar_json(ARQUIVO_PENDENCIAS, pend)


def atualizar_pendencia(idx: int, dados: dict):
    pend = carregar_json(ARQUIVO_PENDENCIAS, pendencias_padrao())
    if 0 <= idx < len(pend):
        pend[idx].update(dados)
        salvar_json(ARQUIVO_PENDENCIAS, pend)


# ==========================================================
# EXCEL
# ==========================================================
def pasta_excel_destino() -> Path:
    p = Path.home() / "OneDrive" / "Documentos" / "ROBO SIEPE"
    p.mkdir(parents=True, exist_ok=True)
    return p


def _get_excel_app():
    try:
        return win32com.client.GetActiveObject("Excel.Application")
    except Exception:
        return None


def _close_workbook_if_open(excel_app, caminho_arquivo: str) -> bool:
    if excel_app is None:
        return False
    try:
        nome_base = os.path.basename(caminho_arquivo)
        for wb in list(excel_app.Workbooks):
            if wb.Name == nome_base:
                wb.Save()
                wb.Close()
                return True
    except Exception:
        pass
    return False


def _reopen_workbook_if_needed(excel_app, caminho_arquivo: str, was_open: bool):
    if not was_open:
        return
    try:
        if excel_app is None:
            excel_app = win32com.client.Dispatch("Excel.Application")
        excel_app.Visible = True
        excel_app.Workbooks.Open(caminho_arquivo)
    except Exception:
        pass


def _abrir_excel_na_aba(caminho_arquivo: str, nome_aba: str) -> None:
    """Abre o Excel sempre na aba do último registro para feedback imediato do usuário."""
    try:
        excel_app = _get_excel_app()
        if excel_app is None:
            excel_app = win32com.client.Dispatch("Excel.Application")
        excel_app.Visible = True

        wb = None
        nome_base = os.path.basename(caminho_arquivo)
        for cand in list(excel_app.Workbooks):
            if str(cand.Name).lower() == nome_base.lower():
                wb = cand
                break
        if wb is None:
            wb = excel_app.Workbooks.Open(caminho_arquivo)

        try:
            ws = wb.Worksheets(nome_aba)
            ws.Activate()
        except Exception:
            pass

        try:
            excel_app.ActiveWindow.ScrollRow = 1
            excel_app.ActiveWindow.ScrollColumn = 1
        except Exception:
            pass
    except Exception:
        pass


def _abrir_ou_criar_planilha(path: Path):
    if path.exists():
        return openpyxl.load_workbook(path)
    wb = Workbook()
    if "Sheet" in wb.sheetnames:
        wb.remove(wb["Sheet"])
    return wb


def _garantir_headers(ws, headers):
    if ws.max_row < 1 or ws.cell(1, 1).value is None:
        ws.append(headers)
        return
    for i, h in enumerate(headers, 1):
        ws.cell(1, i, value=ws.cell(1, i).value or h)


def _concatenar_comentario(celula, novo_texto: str):
    if not novo_texto:
        return
    ts = datetime.now().strftime("%d/%m %H:%M")
    bloco = f"[{ts}] {novo_texto}"
    if celula.comment and celula.comment.text:
        celula.comment = Comment(celula.comment.text + "\n---\n" + bloco, "Megatron")
    else:
        celula.comment = Comment(bloco, "Megatron")


def _atualizar_resumo_semanal(wb, turma_sigla: str):
    aba_diario = f"{turma_sigla} - Diário"
    if aba_diario not in wb.sheetnames:
        return
    ws = wb[aba_diario]
    wsr = wb["Resumo Semanal"] if "Resumo Semanal" in wb.sheetnames else wb.create_sheet("Resumo Semanal")
    _garantir_headers(wsr, ["Turma", "Aluno", "Faltas Semana", "% Presença", "Top5 Flag"])

    datas_cols = [c for c in range(2, ws.max_column + 1) if ws.cell(1, c).value]
    ult_5 = datas_cols[-5:] if len(datas_cols) > 5 else datas_cols
    if not ult_5:
        return

    metricas = []
    for r in range(2, ws.max_row + 1):
        aluno = ws.cell(r, 1).value
        if not aluno:
            continue
        faltas = 0
        presencas = 0
        for c in ult_5:
            v = (ws.cell(r, c).value or "").strip().upper()
            if v == "F":
                faltas += 1
            elif v == "P":
                presencas += 1
        total = faltas + presencas
        pct = round((presencas / total) * 100, 2) if total else 0.0
        metricas.append((aluno, faltas, pct))

    top = {a for a, _, _ in sorted(metricas, key=lambda x: x[1], reverse=True)[:5]}

    # remove linhas antigas da turma
    novas = [
        [wsr.cell(r, c).value for c in range(1, 6)]
        for r in range(2, wsr.max_row + 1)
        if wsr.cell(r, 1).value != turma_sigla
    ]
    wsr.delete_rows(2, max(0, wsr.max_row - 1))
    for linha in novas:
        wsr.append(linha)
    for aluno, faltas, pct in metricas:
        wsr.append([turma_sigla, aluno, faltas, pct, "TOP5" if aluno in top else ""])


def _aplicar_estilo_avaliacao(cel, nivel: int) -> None:
    if nivel not in FILL_AVALIACAO:
        return
    cel.fill = FILL_AVALIACAO[nivel]
    cel.font = FONT_BRANCA if nivel in {1, 5} else FONT_PRETA


def salvar_no_excel(turma_sigla: str, faltosos: List[str], notas_alunos: Dict[str, Dict[str, str]], data_string: str, dia_letivo: bool = True):
    nome_arquivo = pasta_excel_destino() / "Faltas_Megatron.xlsx"
    aba_ind = f"{turma_sigla} - Diário"
    fazer_backup_excel(str(nome_arquivo))
    excel_app = _get_excel_app()
    was_open = _close_workbook_if_open(excel_app, str(nome_arquivo))

    try:
        wb = _abrir_ou_criar_planilha(nome_arquivo)
        ws_log = wb[turma_sigla] if turma_sigla in wb.sheetnames else wb.create_sheet(turma_sigla)
        _garantir_headers(ws_log, ["Data", "Horário", "Turma", "Alunos Faltosos", "Total Faltas", "Observações do Dia"])
        obs_geral = " | ".join([f"{a}: {n['texto']}" for a, n in notas_alunos.items()])
        nomes_faltosos = "NÃO LETIVO" if not dia_letivo else (", ".join(faltosos) if faltosos else "Todos presentes")
        ws_log.append([data_string, datetime.now().strftime("%H:%M:%S"), turma_sigla, nomes_faltosos, len(faltosos) if dia_letivo else 0, obs_geral])

        ws = wb[aba_ind] if aba_ind in wb.sheetnames else wb.create_sheet(aba_ind)
        _garantir_headers(ws, ["Nome do Aluno"])
        ws.freeze_panes = "B2"
        ws["A1"].comment = Comment("Legenda: P=Presente | F=Falta | N=Não letivo", "Megatron")

        cadastrados = {ws.cell(r, 1).value for r in range(2, ws.max_row + 1)}
        for a in banco_alunos.get(turma_sigla, []):
            if a not in cadastrados:
                ws.append([a])

        col_data = None
        for c in range(2, ws.max_column + 1):
            if ws.cell(1, c).value == data_string:
                col_data = c
                break
        if col_data is None:
            col_data = ws.max_column + 1
            ws.cell(1, col_data, value=data_string)

        falt_set = set(faltosos)
        for r in range(2, ws.max_row + 1):
            nome = ws.cell(r, 1).value
            if not nome:
                continue
            status = "" if not dia_letivo else ("F" if nome in falt_set else "P")
            cel = ws.cell(r, col_data, value=status)
            if status == "F":
                cel.fill = FILL_FALTA
            if nome in notas_alunos:
                nota_info = notas_alunos[nome] if isinstance(notas_alunos[nome], dict) else {"texto": str(notas_alunos[nome])}
                texto_nota = str(nota_info.get("texto", "")).strip()
                if texto_nota:
                    _concatenar_comentario(cel, texto_nota)

                nivel_raw = nota_info.get("nivel", 0)
                try:
                    nivel = int(str(nivel_raw).strip() or "0")
                except Exception:
                    nivel = 0

                if 1 <= nivel <= 5:
                    _aplicar_estilo_avaliacao(cel, nivel)
                else:
                    cel.fill = FILL_NOTA

        _atualizar_resumo_semanal(wb, turma_sigla)
        wb.save(nome_arquivo)
        wb.close()
        log_info(f"Excel atualizado para {turma_sigla} ({data_string}).")
    finally:
        _reopen_workbook_if_needed(excel_app, str(nome_arquivo), was_open)
        _abrir_excel_na_aba(str(nome_arquivo), aba_ind)


# ==========================================================
# SIEPE (reuso + otimizações)
# ==========================================================
def _build_driver_visible():
    opcoes = Options()
    opcoes.page_load_strategy = "eager"
    opcoes.add_argument("--start-maximized")
    opcoes.add_argument("--disable-gpu")
    opcoes.add_argument("--no-sandbox")
    opcoes.add_argument("--disable-dev-shm-usage")
    opcoes.add_argument("--remote-allow-origins=*")

    erros = []
    try:
        return webdriver.Chrome(options=opcoes)
    except Exception as exc:
        erros.append(f"padrão: {exc}")

    driver_path = os.environ.get("CHROMEDRIVER") or shutil.which("chromedriver") or shutil.which("chromedriver.exe")
    if driver_path:
        try:
            return webdriver.Chrome(service=Service(driver_path), options=opcoes)
        except Exception as exc:
            erros.append(f"service({driver_path}): {exc}")

    raise RuntimeError("Não foi possível abrir o ChromeDriver. Ajuste CHROMEDRIVER/instalação do Chrome. " + " | ".join(erros))


def _click_js(driver, element):
    driver.execute_script("arguments[0].scrollIntoView({block:'center'});", element)
    time.sleep(0.1)
    driver.execute_script("arguments[0].click();", element)


def _click_locator(driver, locator: Tuple[By, str], timeout: int = 25, tries: int = 6):
    last_exc = None
    for _ in range(tries):
        try:
            el = WebDriverWait(driver, timeout).until(EC.element_to_be_clickable(locator))
            try:
                el.click()
            except (ElementClickInterceptedException, WebDriverException):
                _click_js(driver, el)
            return
        except (StaleElementReferenceException, TimeoutException, WebDriverException) as exc:
            last_exc = exc
            time.sleep(0.35)
    raise last_exc


def _cache_frame(frame_key: str, frame_idx: Optional[int]) -> None:
    FRAME_CACHE[frame_key or "default"] = frame_idx


def _switch_to_frame_key(driver, frame_key: str) -> bool:
    idx = FRAME_CACHE.get(frame_key or "default")
    if idx is None:
        driver.switch_to.default_content()
        return True
    try:
        driver.switch_to.default_content()
        driver.switch_to.frame(idx)
        return True
    except Exception:
        FRAME_CACHE[frame_key or "default"] = None
        return False


def _find_in_any_frame(driver, by, value, timeout=20, frame_key: str = "default"):
    end = time.time() + timeout
    while time.time() < end:
        if _switch_to_frame_key(driver, frame_key):
            try:
                el = driver.find_element(by, value)
                return el, FRAME_CACHE.get(frame_key)
            except Exception:
                pass

        driver.switch_to.default_content()
        try:
            el = driver.find_element(by, value)
            _cache_frame(frame_key, None)
            return el, None
        except Exception:
            pass

        for i, _ in enumerate(driver.find_elements(By.TAG_NAME, "iframe")):
            try:
                driver.switch_to.default_content()
                driver.switch_to.frame(i)
                el = driver.find_element(by, value)
                _cache_frame(frame_key, i)
                return el, i
            except Exception:
                continue
        time.sleep(0.15)
    raise TimeoutException(f"Elemento não encontrado: {by}={value}")


def _find_all_in_any_frame(driver, by, value, timeout=12, frame_key: str = "default"):
    end = time.time() + timeout
    while time.time() < end:
        if _switch_to_frame_key(driver, frame_key):
            try:
                els = driver.find_elements(by, value)
                if els:
                    return els, FRAME_CACHE.get(frame_key)
            except Exception:
                pass

        try:
            driver.switch_to.default_content()
            els = driver.find_elements(by, value)
            if els:
                _cache_frame(frame_key, None)
                return els, None
        except Exception:
            pass

        try:
            iframes = driver.find_elements(By.TAG_NAME, "iframe")
        except Exception:
            iframes = []
        for i, _ in enumerate(iframes):
            try:
                driver.switch_to.default_content()
                driver.switch_to.frame(i)
                els = driver.find_elements(by, value)
                if els:
                    _cache_frame(frame_key, i)
                    return els, i
            except Exception:
                continue
        time.sleep(0.15)
    return [], None


def _try_switch_any_frame_with(driver, by: By, value: str, frame_key: str = "default") -> bool:
    if _switch_to_frame_key(driver, frame_key):
        try:
            if driver.find_elements(by, value):
                return True
        except Exception:
            pass

    driver.switch_to.default_content()
    if driver.find_elements(by, value):
        _cache_frame(frame_key, None)
        return True
    for i, _ in enumerate(driver.find_elements(By.TAG_NAME, "iframe")):
        try:
            driver.switch_to.default_content()
            driver.switch_to.frame(i)
            if driver.find_elements(by, value):
                _cache_frame(frame_key, i)
                return True
        except Exception:
            continue
    driver.switch_to.default_content()
    return False


def _login_educadores(driver, apelido: str, senha: str):
    """Login resiliente no SIEPE com validação pós-login e fallback de locators."""
    driver.get(SIEPE_URL)
    radar = WebDriverWait(driver, 30)

    login_inp = radar.until(EC.element_to_be_clickable((By.ID, "login")))
    login_inp.clear()
    login_inp.send_keys(apelido)

    try:
        senha_inp = driver.find_element(By.ID, "senha")
    except Exception:
        senha_inp, _ = _find_in_any_frame(driver, By.ID, "senha", timeout=8, frame_key="default")
    senha_inp.clear()
    senha_inp.send_keys(senha)
    senha_inp.send_keys(Keys.RETURN)

    erros_login = [
        (By.XPATH, "//*[contains(translate(.,'INVALIDOCREDENCIAIS','invalidocredenciais'),'inválid') or contains(translate(.,'INVALIDOCREDENCIAIS','invalidocredenciais'),'credencia') or contains(translate(.,'INVALIDOCREDENCIAIS','invalidocredenciais'),'senha')]")
    ]

    fim = time.time() + 25
    while time.time() < fim:
        try:
            _find_in_any_frame(driver, SEL_EDUCADORES[0], SEL_EDUCADORES[1], timeout=2, frame_key="default")
            return
        except Exception:
            pass
        for loc in erros_login:
            try:
                err, _ = _find_in_any_frame(driver, loc[0], loc[1], timeout=1, frame_key="default")
                msg = (err.text or "").strip()
                raise RuntimeError(f"Login SIEPE com apelido/senha inválidos ou bloqueado: {msg or 'verifique credenciais/captcha'}")
            except RuntimeError:
                raise
            except Exception:
                continue
        time.sleep(0.2)

    raise RuntimeError("Login no SIEPE não confirmou acesso ao menu de Educadores.")


def _sincronizar_alunos_da_turma_no_siepe(driver, sigla_turma: str):
    if not sigla_turma:
        return

    fim = time.time() + 12
    links = []
    while time.time() < fim:
        links = driver.find_elements(By.XPATH, "//a[contains(@onclick, 'mostrarInformaçõesAluno')]")
        if len(links) >= 5:
            break
        time.sleep(0.25)

    if not links:
        links, _ = _find_all_in_any_frame(driver, By.XPATH, "//a[contains(@onclick, 'mostrarInformaçõesAluno')]", timeout=6, frame_key="frequencia")

    nomes: List[str] = []
    seen = set()

    # caminho principal: coleta via JS para evitar stale element durante render dinâmico do SIEPE
    try:
        nomes_js = driver.execute_script(
            """
            return Array.from(document.querySelectorAll("a[onclick*='mostrarInformaçõesAluno']"))
              .map(a => (a.innerText || a.textContent || '').trim())
              .filter(Boolean);
            """
        ) or []
        for n in nomes_js:
            nome = str(n).strip().upper()
            if nome and nome not in seen:
                seen.add(nome)
                nomes.append(nome)
    except Exception:
        pass

    # fallback linear: recarrega o texto de cada elemento com tolerância a stale
    if len(nomes) < 5:
        for i in range(len(links)):
            try:
                itens = driver.find_elements(By.XPATH, "//a[contains(@onclick, 'mostrarInformaçõesAluno')]")
                if i >= len(itens):
                    continue
                nome = (itens[i].text or "").strip().upper()
                if nome and nome not in seen:
                    seen.add(nome)
                    nomes.append(nome)
            except StaleElementReferenceException:
                continue
            except Exception:
                continue

    if len(nomes) < 5:
        log_warn(f"Poucos alunos detectados para {sigla_turma} ({len(nomes)}). Não atualizei banco.")
        return

    banco_alunos[sigla_turma] = nomes
    recarregar_indices_alunos()
    salvar_json(str(ARQUIVO_BANCO_ALUNOS), banco_alunos)
    log_info(f"Base de alunos sincronizada do SIEPE para {sigla_turma}: {len(nomes)} alunos.")

def carregar_credenciais():
    if CRED_FILE.exists():
        linhas = CRED_FILE.read_text(encoding="utf-8").splitlines()
        if len(linhas) >= 2:
            return linhas[0], linhas[1]
    return "", ""


def pedir_credenciais_interativas() -> Tuple[str, str]:
    apelido = input("Apelido: ").strip()
    senha = getpass.getpass("Senha: ")
    if apelido and senha:
        CRED_FILE.write_text(f"{apelido}\n{senha}", encoding="utf-8")
    return apelido, senha


def _eh_erro_login_siepe(exc: Exception) -> bool:
    tx = normalizar_texto(str(exc))
    chaves = ["LOGIN SIEPE", "INVALID", "CREDENCIA", "CAPTCHA", "BLOQUEADO", "BLOQUEIO", "APELIDO/SENHA"]
    return any(k in tx for k in chaves)


def _login_educadores_com_retry(driver, apelido: str, senha: str) -> Tuple[str, str]:
    while True:
        try:
            _login_educadores(driver, apelido, senha)
            return apelido, senha
        except Exception as exc:
            if not _eh_erro_login_siepe(exc):
                raise
            acao = tui_select(
                "Login SIEPE",
                f"Login falhou: {exc}\nDeseja reinformar apelido/senha?",
                [("retry", "Reinformar credenciais"), ("cancel", "Cancelar")],
                "retry",
            )
            if acao in {None, NAV_BACK, NAV_MENU, "cancel"}:
                raise RuntimeError("Login cancelado pelo usuário.") from exc
            apelido, senha = pedir_credenciais_interativas()
            if not apelido or not senha:
                raise RuntimeError("Credenciais não informadas.") from exc
            try:
                driver.get(SIEPE_URL)
            except Exception:
                pass


def _set_data_frequencia(driver, data_foco: datetime):
    data_visivel = data_foco.strftime("%d/%m/%Y")
    inp, _ = _find_in_any_frame(driver, By.ID, "txtDataFrequenciaDisciplina", timeout=10, frame_key="frequencia")
    inp.send_keys(Keys.CONTROL, "a")
    inp.send_keys(data_visivel)
    inp.send_keys(Keys.TAB)

    try:
        WebDriverWait(driver, 8).until(lambda d: len(d.find_elements(By.XPATH, "//input[contains(@id, 'FNJ') and not(@disabled)]")) > 0)
        return
    except Exception:
        pass

    try:
        driver.execute_script(
            """
            var dt = arguments[0];
            var v = document.getElementById('txtDataFrequenciaDisciplina');
            if (v) { v.value = dt; v.dispatchEvent(new Event('change')); }
            var h = document.getElementById('hdnDataFrequenciaDisciplina');
            if (h) { h.value = dt; }
            """,
            data_visivel,
        )
    except Exception:
        pass


def _coletar_checkboxes_frequencia(driver):
    caixas, idx = _find_all_in_any_frame(driver, By.XPATH, "//input[contains(@id, 'FNJ')]", timeout=10, frame_key="frequencia")
    if idx is not None:
        _cache_frame("frequencia", idx)
    return caixas


def _coletar_mapa_aluno_id(driver) -> Dict[str, str]:
    links, _ = _find_all_in_any_frame(driver, By.XPATH, "//a[contains(@onclick, 'mostrarInformaçõesAluno')]", timeout=8, frame_key="frequencia")
    mapa: Dict[str, str] = {}
    for l in links:
        nome = normalizar_texto((l.text or "").strip())
        oc = l.get_attribute("onclick") or ""
        if nome and "'" in oc:
            mapa[nome] = oc.split("'")[1]
    return mapa


def _marcar_frequencia(driver, modo_lancamento: str, lista_digitada: List[str], caixas, mapa_aluno_id: Dict[str, str]):
    if modo_lancamento == "P":
        for chk in caixas:
            if not chk.is_selected():
                _click_js(driver, chk)

    for nome in lista_digitada or []:
        nn = normalizar_texto(nome)
        aluno_id = mapa_aluno_id.get(nn)
        if not aluno_id:
            continue
        chks = driver.find_elements(By.XPATH, f"//input[@value='{aluno_id}' and contains(@id, 'FNJ')]")
        for chk in chks:
            if modo_lancamento == "P" and chk.is_selected():
                _click_js(driver, chk)
            elif modo_lancamento != "P" and not chk.is_selected():
                _click_js(driver, chk)


def _confirmar_gravacao(driver):
    sinais = [
        (By.XPATH, "//*[contains(translate(.,'SALVOGRAVADOSUCESSO','salvogravadosucesso'),'sucesso') or contains(translate(.,'SALVOGRAVADOSUCESSO','salvogravadosucesso'),'gravado') or contains(translate(.,'SALVOGRAVADOSUCESSO','salvogravadosucesso'),'salvo') ]"),
        (By.ID, "chkConclusaoFrequenciaDiaria"),
    ]
    fim = time.time() + 12
    while time.time() < fim:
        for loc in sinais:
            try:
                el, _ = _find_in_any_frame(driver, loc[0], loc[1], timeout=1.3, frame_key="frequencia")
                if loc[0] == By.ID:
                    if el.is_selected():
                        return
                else:
                    if (el.text or "").strip():
                        return
            except Exception:
                continue
        time.sleep(0.2)
    raise RuntimeError("Não consegui confirmar gravação da frequência no SIEPE.")


def _navegar_ate_frequencia(driver, texto_busca_siepe: str):
    _abrir_diario_classe(driver)
    _selecionar_turma_like_old(driver, texto_busca_siepe)
    _abrir_sanfona_frequencia_like_old(driver)


def executar_siepe(meu_apelido, minha_senha, texto_busca_siepe, modo_lancamento, lista_digitada, data_foco: datetime, force_fail=False, sigla_turma: str = "") -> bool:
    if force_fail:
        raise RuntimeError("Falha simulada para teste de pendências.")

    driver = _build_driver_visible()
    try:
        for tentativa in range(1, 4):
            try:
                meu_apelido, minha_senha = _login_educadores_com_retry(driver, meu_apelido, minha_senha)
                _navegar_ate_frequencia(driver, texto_busca_siepe)
                _set_data_frequencia(driver, data_foco)

                caixas = _coletar_checkboxes_frequencia(driver)
                if not caixas:
                    return False

                mapa = _coletar_mapa_aluno_id(driver)
                if sigla_turma:
                    _sincronizar_alunos_da_turma_no_siepe(driver, sigla_turma)
                _marcar_frequencia(driver, modo_lancamento, lista_digitada, caixas, mapa)

                try:
                    c, _ = _find_in_any_frame(driver, By.ID, "chkConclusaoFrequenciaDiaria", timeout=4, frame_key="frequencia")
                    if not c.is_selected():
                        _click_js(driver, c)
                except Exception:
                    pass

                g, _ = _find_in_any_frame(driver, By.ID, "btnGravarFrequenciaDisciplina", timeout=15, frame_key="frequencia")
                _click_js(driver, g)
                _confirmar_gravacao(driver)
                return True
            except (TimeoutException, WebDriverException, RuntimeError) as exc:
                if tentativa == 3:
                    raise RuntimeError("Falha de automação no SIEPE. Site instável.") from exc
                try:
                    driver.get(SIEPE_URL)
                except Exception:
                    pass
                time.sleep(0.8)
    finally:
        try:
            driver.quit()
        except Exception:
            pass


# ==========================================================
# IMPORTAÇÃO OTIMIZADA COM PREVIEW
# ==========================================================
def normalizar_horario(s: str) -> str:
    s = (s or "").strip()
    m = re.findall(r"(\d{2}:\d{2})", s)
    return f"{m[0]} - {m[1]}" if len(m) >= 2 else s


def inferir_turma_do_quadro(item: Dict[str, Any]) -> str:
    """Mantido por compatibilidade, mas agora conservador (maioria clara em textos úteis)."""
    texto = normalizar_texto(f"{item.get('label','')} {item.get('titulo','')}")
    turmas = _extrair_turmas_no_texto(texto)
    if len(turmas) == 1:
        return turmas[0]

    grade = item.get("grade") or {}
    cont: Dict[str, int] = {}
    total = 0
    for dia in DIAS_COL:
        for _h, val in (grade.get(dia) or {}).items():
            cands = _extrair_turmas_no_texto(str(val or ""))
            if len(cands) == 1:
                t = cands[0]
                total += 1
                cont[t] = cont.get(t, 0) + 1

    if not cont:
        return ""
    top_turma, top_n = sorted(cont.items(), key=lambda x: x[1], reverse=True)[0]
    second_n = sorted(cont.values(), reverse=True)[1] if len(cont) > 1 else 0
    if total >= 3 and (top_n / total) >= 0.65 and (top_n - second_n) >= 2:
        return top_turma
    return ""



class ProfessorCompleter(Completer):
    def __init__(self, professores: List[str]):
        self.professores = [(p, normalizar_texto(p)) for p in professores]

    def get_completions(self, document, complete_event):
        q = normalizar_texto(document.text)
        if not q:
            return
        candidatos = [(p, pnorm) for p, pnorm in self.professores if _match_priority(pnorm, q)[0] < 9]
        candidatos.sort(key=lambda x: _match_priority(x[1], q))
        for p, _ in candidatos:
            yield Completion(p, start_position=-len(document.text))


def atualizar_memoria_professores(coletados: List[Dict[str, Any]]) -> List[str]:
    profs = set()
    for item in coletados:
        matriz = item.get("matriz") or {}
        for _sigla, info in matriz.items():
            p = (info.get("professor") or "").strip()
            if p and "NÃO ATRIBU" not in p.upper() and "NAO ATRIBU" not in p.upper():
                profs.add(p)
    lista = sorted(profs, key=lambda x: x.upper())
    if lista:
        MEM_PROFS.write_text(json.dumps(lista, ensure_ascii=False, indent=4), encoding="utf-8")
    return lista


def ajustar_horarios_padrao_dinamico(horarios_reais: List[str]) -> None:
    global HORARIOS_PADRAO

    def key_time(x: str):
        m = re.search(r"(\d{2}):(\d{2})", x)
        return (int(m.group(1)), int(m.group(2))) if m else (99, 99)

    base = []
    for h in horarios_reais:
        hn = normalizar_horario(h)
        if hn and "-" in hn and hn not in base:
            base.append(hn)

    base.sort(key=key_time)
    base = base[:7]

    if len(base) < 7:
        for h in HORARIOS_PADRAO:
            if h not in base:
                base.append(h)
            if len(base) == 7:
                break

    HORARIOS_PADRAO = base


PERMITIDAS_TURMAS = {"6A", "6B", "7A", "7B", "8A", "8B", "9A", "9B"}


def _normalizar_professor_chave(txt: str) -> str:
    base = normalizar_texto(txt)
    base = re.sub(r"[^A-Z0-9 ]+", " ", base)
    base = re.sub(r"\s+", " ", base).strip()
    return base


def _extrair_turmas_no_texto(texto: str, modo_slot: bool = False) -> List[str]:
    """Extrai siglas válidas 6A..9B com heurística conservadora para evitar falso positivo."""
    norm = normalizar_texto(texto or "")
    if not norm:
        return []

    # Em célula/slot, não usar compostos ambíguos (6A/7A, 6A - ED.FIS etc.).
    # Em label/título (modo_slot=False) o '/' é permitido.
    if modo_slot:
        if "/" in norm:
            return []
        if re.search(r"\b[6-9][AB]\s*-\s*[A-Z]", norm):
            return []

    tokens = re.findall(r"(?<![A-Z0-9])([6-9][AB])(?![A-Z0-9])", norm)
    out: List[str] = []
    for t in tokens:
        if t in PERMITIDAS_TURMAS and t not in out:
            out.append(t)
    return out



def _log_estrutura(evento: str, payload: Dict[str, Any]) -> None:
    try:
        log_info(f"[STRUCT] {evento} | {json.dumps(payload, ensure_ascii=False)}")
    except Exception:
        log_info(f"[STRUCT] {evento} | {payload}")


def _resolver_turma_slot(item: Dict[str, Any], dia: str, horario: str, valor_celula: str, mapeamento_manual: Dict[str, str]) -> str:
    """Não inferir por disciplina da célula; usa só meta da turma do quadro."""
    contexto = item.get("contexto_turma") or {}
    meta_sigla = str((contexto.get("turma_sigla") or "")).strip().upper()
    if meta_sigla in PERMITIDAS_TURMAS:
        return meta_sigla
    return ""

def _resolver_turma_item(item: Dict[str, Any], mapeamento_manual: Dict[str, str]) -> str:
    """
    Resolve turma do quadro.
    Regra: meta do detalhe (contexto_turma) tem prioridade sobre mapeamento id.
    """
    qid = str(item.get("id", "")).strip()
    lbl = str(item.get("label", "")).strip()
    titulo = str(item.get("titulo", "")).strip()
    contexto_turma = item.get("contexto_turma") or {}

    meta_sigla = str((contexto_turma.get("turma_sigla") or "")).strip().upper()
    if meta_sigla in PERMITIDAS_TURMAS:
        if qid:
            mapped = str((mapeamento_manual.get(f"id:{qid}") or "")).strip().upper()
            if mapped in PERMITIDAS_TURMAS and mapped != meta_sigla:
                log_warn(f"[MAP] Conflito id:{qid}: mapeamento={mapped} vs meta={meta_sigla}. Usando META.")
        return meta_sigla

    if qid:
        turma_id = (mapeamento_manual.get(f"id:{qid}") or "").strip().upper()
        if turma_id in PERMITIDAS_TURMAS:
            return turma_id

    cand_txt = _extrair_turmas_no_texto(f"{lbl} {titulo}", modo_slot=False)
    if len(cand_txt) == 1:
        return cand_txt[0]
    if len(cand_txt) > 1:
        return ""

    return ""

def montar_grade_professor(coletados: List[Dict[str, Any]], professor: str, mapeamento_manual: Dict[str, str]) -> Tuple[Dict[str, Dict[str, str]], List[str], List[str]]:
    """Monta grade por professor priorizando slots realmente atribuídos ao docente."""
    prof_key = _normalizar_professor_chave(professor)

    def prof_match(prof_cell: str) -> bool:
        alvo = _normalizar_professor_chave(prof_cell)
        if not alvo:
            return False
        return alvo == prof_key or prof_key in alvo or alvo in prof_key

    all_times: List[str] = []
    dia_idx = {"SEG": "0", "TER": "1", "QUA": "2", "QUI": "3", "SEX": "4"}
    grade_final: Dict[str, Dict[str, str]] = {str(i): {} for i in range(5)}
    conflitos: List[str] = []

    por_quadro: Dict[str, Dict[str, int]] = {}

    for item in coletados:
        grade = item.get("grade") or {}
        matriz = item.get("matriz") or {}
        qid = str(item.get("id", "")).strip()
        lbl = str(item.get("label", "")).strip()

        sigla_to_prof = {}
        for sigla, info in matriz.items():
            p = (info.get("professor") or "").strip()
            if p:
                sigla_to_prof[sigla.strip()] = p

        turma_quadro = _resolver_turma_item(item, mapeamento_manual)
        _log_estrutura("quadro_importado", {
            "id": qid,
            "label": lbl,
            "turma_resolvida": turma_quadro or "",
            "turma_meta": (item.get("contexto_turma") or {}).get("turma_sigla", ""),
        })

        cont_quadro: Dict[str, int] = {}
        total_quadro = 0

        for dia in DIAS_COL:
            for h, sigla in (grade.get(dia) or {}).items():
                hh = normalizar_horario(h)
                sig = (sigla or "").strip()
                if not hh or not sig:
                    continue

                prof_cell = sigla_to_prof.get(sig, "")
                if not prof_match(prof_cell):
                    continue

                if hh not in all_times:
                    all_times.append(hh)

                turma_slot = _resolver_turma_slot(item, dia, hh, sig, mapeamento_manual)
                if not turma_slot and turma_quadro:
                    turma_slot = turma_quadro

                if turma_slot not in PERMITIDAS_TURMAS:
                    continue

                total_quadro += 1
                cont_quadro[turma_slot] = cont_quadro.get(turma_slot, 0) + 1
                _log_estrutura("slot_atribuido", {
                    "id": qid,
                    "label": lbl,
                    "dia": dia,
                    "horario": hh,
                    "valor_original": sig,
                    "turma_final": turma_slot,
                })

                d = dia_idx[dia]
                atual = grade_final[d].get(hh, "AULA ATIVIDADE")
                if atual == "AULA ATIVIDADE":
                    grade_final[d][hh] = turma_slot
                elif turma_slot not in atual.split(" | "):
                    grade_final[d][hh] = f"{atual} | {turma_slot}"
                    conflitos.append(f"{dia} {hh}: {atual} + {turma_slot}")

        if total_quadro > 0:
            por_quadro[qid or lbl or f"idx:{len(por_quadro)+1}"] = cont_quadro
            top_turma, top_n = sorted(cont_quadro.items(), key=lambda x: x[1], reverse=True)[0]
            dominancia = top_n / max(1, total_quadro)
            if total_quadro >= 5 and dominancia >= 0.9:
                log_warn(f"Anomalia potencial no quadro {qid or lbl}: {top_turma} domina {top_n}/{total_quadro} slots.")

    def key_time(x: str):
        m = re.search(r"(\d{2}):(\d{2})", x)
        return (int(m.group(1)), int(m.group(2))) if m else (99, 99)

    all_times.sort(key=key_time)
    return grade_final, all_times, conflitos



def _abrir_quadro_horarios(driver):
    _click_locator(driver, SEL_EDUCADORES, timeout=30, tries=8)
    candidates = [
        (By.ID, "lnkQuadroHorario53238446"),
        (By.XPATH, "//*[contains(@class,'ac-icon-quadro-horarios')]"),
        (By.XPATH, "//button[contains(.,'Quadro de Horários')]"),
        (By.XPATH, "//h1[contains(.,'Quadro de Horários')]/ancestor::*[1]"),
    ]
    clicked = False
    for loc in candidates:
        try:
            if not _try_switch_any_frame_with(driver, loc[0], loc[1], frame_key="quadro"):
                driver.switch_to.default_content()
            if driver.find_elements(loc[0], loc[1]):
                _click_locator(driver, loc, timeout=15, tries=3)
                clicked = True
                break
        except Exception:
            continue
    if not clicked:
        driver.switch_to.default_content()
    if not _try_switch_any_frame_with(driver, By.ID, ID_CONTAINER_TURMAS, frame_key="quadro"):
        driver.switch_to.default_content()
    WebDriverWait(driver, 30).until(lambda d: len(d.find_elements(By.ID, ID_CONTAINER_TURMAS)) > 0)


def _esperar_turmas_carregarem(driver, timeout: int = 30):
    if not _try_switch_any_frame_with(driver, By.ID, ID_CONTAINER_TURMAS, frame_key="quadro"):
        driver.switch_to.default_content()

    WebDriverWait(driver, timeout).until(lambda d: len(d.find_elements(By.ID, ID_CONTAINER_TURMAS)) > 0)

    fim = time.time() + timeout
    ultimo = -1
    estavel = 0
    while time.time() < fim:
        total = len(driver.find_elements(By.XPATH, XPATH_LINKS_DETALHE))
        if total > 0 and total == ultimo:
            estavel += 1
            if estavel >= 3:
                return
        else:
            estavel = 0
        ultimo = total
        time.sleep(0.25)
    raise TimeoutException("Lista de turmas não estabilizou no quadro de horários.")



def _listar_ids_quadros(driver) -> List[Tuple[str, str]]:
    _esperar_turmas_carregarem(driver, timeout=35)
    links = driver.find_elements(By.XPATH, XPATH_LINKS_DETALHE)
    out = []
    for a in links:
        blob = f"{a.get_attribute('href') or ''} {a.get_attribute('onclick') or ''}"
        txt = (a.text or "").strip()
        m = re.search(r"detalheQuadroDeHorario\((\d+)\)", blob)
        if m:
            out.append((m.group(1), txt if txt else f"Quadro {m.group(1)}"))
    seen, uniq = set(), []
    for qid, label in out:
        if qid not in seen:
            uniq.append((qid, label))
            seen.add(qid)
    return uniq


def _abrir_detalhe_por_id_js(driver, quadro_id: str):
    if not _try_switch_any_frame_with(driver, By.ID, "hdnIdTurma", frame_key="quadro"):
        driver.switch_to.default_content()
    driver.execute_script("detalheQuadroDeHorario(arguments[0]);", int(quadro_id))
    if not _try_switch_any_frame_with(driver, By.ID, "btnVoltarDetalhe", frame_key="quadro"):
        driver.switch_to.default_content()
    WebDriverWait(driver, 30).until(EC.presence_of_element_located((By.ID, "btnVoltarDetalhe")))
    if not _try_switch_any_frame_with(driver, By.XPATH, XPATH_ALL_TABELAS_HORARIO, frame_key="quadro"):
        driver.switch_to.default_content()
    WebDriverWait(driver, 30).until(lambda d: len(d.find_elements(By.XPATH, XPATH_ALL_TABELAS_HORARIO)) >= 1)


def _extrair_contexto_turma_detalhe(driver) -> Dict[str, str]:
    """Extrai metadados da turma no detalhe (linha Turma:/Ano/Turno etc.)."""
    txt = ""
    try:
        txt = (driver.find_element(By.TAG_NAME, "body").text or "").strip()
    except Exception:
        txt = normalizar_texto(driver.page_source or "")

    n = normalizar_texto(txt)
    turma_linha = ""
    sigla = ""
    try:
        m = re.search(r"TURMA\s*:\s*([^\n\r]+)", txt, flags=re.IGNORECASE)
        if m:
            turma_linha = m.group(1).strip()
        else:
            m2 = re.search(r"TURMA\s*:\s*([^\n\r]+)", n)
            if m2:
                turma_linha = m2.group(1).strip()
    except Exception:
        pass

    if turma_linha:
        cands = _extrair_turmas_no_texto(turma_linha, modo_slot=False)
        if len(cands) == 1:
            sigla = cands[0]

    if not sigla:
        cands2 = _extrair_turmas_no_texto(txt, modo_slot=False)
        if len(cands2) == 1:
            sigla = cands2[0]

    return {"turma_linha": turma_linha, "turma_sigla": sigla}



def _title_of_table(table_el) -> str:
    try:
        return (table_el.find_element(By.XPATH, ".//thead//th[contains(@class,'Title')]").text or "").strip()
    except Exception:
        return ""


def _extrair_quadros_e_matriz_do_detalhe(driver):
    tables = driver.find_elements(By.XPATH, XPATH_ALL_TABELAS_HORARIO)
    if not tables:
        raise RuntimeError("Não encontrei tabelas ListagemPadrao Horario no detalhe.")

    quadro_table = None
    matriz_table = None
    for t in tables:
        title = _title_of_table(t).upper()
        if "MATRIZ CURRICULAR" in title:
            matriz_table = t
        elif "QUADRO DE HOR" in title:
            quadro_table = t

    if quadro_table is None:
        for t in tables:
            head = " ".join([x.text.strip().upper() for x in t.find_elements(By.XPATH, ".//thead//tr[2]/th")])
            if "HOR" in head and "SEG" in head:
                quadro_table = t
                break

    if quadro_table is None:
        raise RuntimeError("Não achei a tabela do Quadro de Horários.")

    titulo = _title_of_table(quadro_table) or "Quadro de Horários"

    ths = quadro_table.find_elements(By.XPATH, ".//thead/tr[2]/th")
    headers = [(th.text or "").strip().upper() for th in ths]
    base_dias = {"SEG", "TER", "QUA", "QUI", "SEX"}
    if len(base_dias.intersection(set(headers))) < 3:
        _salvar_diagnostico_coleta(driver, "cabecalho_quadro_invalido")
        raise RuntimeError(f"Cabeçalho inesperado no quadro de horários: {headers}")

    col_to_day = {}
    for cidx, h in enumerate(headers):
        if h in DIAS_HEADERS_ACCEPT:
            col_to_day[cidx] = h

    dados_grade: Dict[str, Dict[str, str]] = {d: {} for d in DIAS_HEADERS_ACCEPT}
    rows = quadro_table.find_elements(By.XPATH, ".//tbody/tr")
    for r in rows:
        try:
            hcell = r.find_element(By.XPATH, "./th")
        except Exception:
            continue
        horario = normalizar_horario(hcell.text)
        tds = r.find_elements(By.XPATH, "./td")
        for head_idx, dia in col_to_day.items():
            td_i = head_idx - 1
            if 0 <= td_i < len(tds):
                val = (tds[td_i].text or "").strip()
                if val:
                    dados_grade[dia][horario] = val

    matriz = {}
    if matriz_table is not None:
        for mr in matriz_table.find_elements(By.XPATH, ".//tbody/tr"):
            tds = mr.find_elements(By.XPATH, "./td")
            if len(tds) < 4:
                continue
            sigla = (tds[0].text or "").strip()
            comp = (tds[1].text or "").strip()
            prof = (tds[2].text or "").strip()
            aulas_txt = (tds[3].text or "").strip()
            if not sigla or "TOTAL" in sigla.upper():
                continue
            aulas = int(re.sub(r"\D+", "", aulas_txt) or "0")
            matriz[sigla] = {"componente": comp, "professor": prof, "aulas": aulas}

    contexto = _extrair_contexto_turma_detalhe(driver)
    return titulo, dados_grade, matriz, contexto


def _voltar_para_lista_turmas(driver):
    if not _try_switch_any_frame_with(driver, By.ID, "btnVoltarDetalhe", frame_key="quadro"):
        driver.switch_to.default_content()
    _click_locator(driver, (By.ID, "btnVoltarDetalhe"), timeout=20, tries=6)
    if not _try_switch_any_frame_with(driver, By.ID, ID_CONTAINER_TURMAS, frame_key="quadro"):
        driver.switch_to.default_content()
    WebDriverWait(driver, 30).until(lambda d: len(d.find_elements(By.ID, ID_CONTAINER_TURMAS)) > 0)
    _esperar_turmas_carregarem(driver, timeout=30)


def _xpath_literal(valor: str) -> str:
    """Literal seguro para xpath contains, tolerante a aspas."""
    safe = (valor or "").replace("'", " ").replace('"', " ")
    safe = re.sub(r"\s+", " ", safe).strip()
    return f"'{safe}'"


def _salvar_diagnostico_coleta(driver, etapa: str) -> None:
    try:
        dump = Path(f"debug_coleta_{etapa}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.html")
        dump.write_text(driver.page_source, encoding="utf-8")
        log_warn(f"Diagnóstico salvo em {dump.name}")
    except Exception:
        pass


def _selecionar_turma_like_old(driver, texto_busca_siepe: str):
    btn, idx = _find_in_any_frame(driver, By.XPATH, "//span[contains(text(), 'Pesquisar')]", timeout=25, frame_key="diario")
    if idx is not None:
        _cache_frame("diario", idx)
    _click_js(driver, btn)

    alvo = normalizar_texto(texto_busca_siepe)
    sig_alvo = _inferir_sigla_turma_por_texto(texto_busca_siepe)
    linhas, _ = _find_all_in_any_frame(driver, By.XPATH, "//tr[td]", timeout=12, frame_key="diario")
    if not linhas:
        _salvar_diagnostico_coleta(driver, "diario_sem_linhas")
        raise RuntimeError("Lista de turmas vazia no Diário de Classe.")

    best = None
    score_best = -1.0
    candidatos: List[Tuple[float, str]] = []

    for tr in linhas:
        try:
            tds = tr.find_elements(By.XPATH, "./td")
            if not tds:
                continue
            linha_txt = " ".join([(td.text or "").strip() for td in tds if (td.text or "").strip()])
            if not linha_txt:
                continue
            txt_norm = normalizar_texto(linha_txt)
            if any(k in txt_norm for k in ["PESQUISAR", "TOTAL", "FILTRO", "NENHUM REGISTRO", "AÇÃO", "ACAO"]):
                continue

            sig_linha = _inferir_sigla_turma_por_texto(linha_txt)

            base = 0.0
            if sig_alvo and sig_linha and sig_alvo == sig_linha:
                base += 1.0
            if sig_alvo and sig_alvo in txt_norm:
                base += 0.7
            if sig_alvo:
                ano = sig_alvo[0]
                turma = sig_alvo[1]
                if re.search(rf"\b{ano}\s*(?:O|º)?\s*ANO\s*{turma}\b", txt_norm):
                    base += 0.9
            if alvo and alvo in txt_norm:
                base += 0.6

            import difflib
            ratio = difflib.SequenceMatcher(None, alvo, txt_norm).ratio()
            score = base + ratio
            candidatos.append((score, linha_txt[:180]))

            if score > score_best:
                score_best = score
                best = tr
        except Exception:
            continue

    # fallback linear/unificado: se existir alvo de sigla, tenta clicar por xpath antes de falhar
    if (best is None or score_best < 0.65) and sig_alvo:
        ano = sig_alvo[0]
        turma = sig_alvo[1]
        xpaths = [
            f"//tr[td][.//*[contains(normalize-space(.), '{sig_alvo}')]]",
            f"//tr[td][contains(normalize-space(.), '{ano}º ANO {turma}') or contains(normalize-space(.), '{ano}o ANO {turma}') or contains(normalize-space(.), '{ano} ANO {turma}')]",
        ]
        for xp in xpaths:
            try:
                linha, _ = _find_in_any_frame(driver, By.XPATH, xp, timeout=2, frame_key="diario")
                if linha:
                    best = linha
                    score_best = max(score_best, 0.70)
                    break
            except Exception:
                continue

    if best is None or score_best < 0.65:
        if candidatos:
            top = sorted(candidatos, key=lambda x: x[0], reverse=True)[:5]
            log_warn("Top candidatos no diário: " + " | ".join([f"{s:.2f}:{t}" for s, t in top]))
        _salvar_diagnostico_coleta(driver, "falha_selecionar_turma_score")
        raise RuntimeError(f"Não consegui selecionar turma no diário: {texto_busca_siepe} (score {score_best:.2f})")

    _click_js(driver, best)
    WebDriverWait(driver, 20).until(
        lambda d: _try_switch_any_frame_with(d, By.XPATH, "//*[contains(@onclick,'divFrequenciaDiaria')]", frame_key="frequencia")
    )

def _abrir_sanfona_frequencia_like_old(driver):
    sanfona, idx = _find_in_any_frame(driver, By.XPATH, "//*[contains(@onclick,'divFrequenciaDiaria')]", timeout=12, frame_key="frequencia")
    if idx is not None:
        _cache_frame("frequencia", idx)
    _click_js(driver, sanfona)
    WebDriverWait(driver, 20).until(
        lambda d: len(d.find_elements(By.ID, "txtDataFrequenciaDisciplina")) > 0
    )


def _abrir_diario_classe(driver):
    radar = WebDriverWait(driver, 25)
    try:
        radar.until(EC.element_to_be_clickable((By.ID, "mnSupEdu"))).click()
    except Exception:
        mn, _ = _find_in_any_frame(driver, By.ID, "mnSupEdu", timeout=12, frame_key="default")
        _click_js(driver, mn)
    try:
        radar.until(EC.element_to_be_clickable((By.CLASS_NAME, "ac-icon-diario-classe"))).click()
    except Exception:
        dc, idx = _find_in_any_frame(driver, By.CLASS_NAME, "ac-icon-diario-classe", timeout=12, frame_key="diario")
        if idx is not None:
            _cache_frame("diario", idx)
        _click_js(driver, dc)
    _find_in_any_frame(driver, By.XPATH, "//span[contains(text(), 'Pesquisar')]", timeout=20, frame_key="diario")


def _listar_turmas_no_diario(driver) -> List[str]:
    btn, _ = _find_in_any_frame(driver, By.XPATH, "//span[contains(text(), 'Pesquisar')]", timeout=20, frame_key="diario")
    _click_js(driver, btn)

    fim = time.time() + 20
    linhas = []
    while time.time() < fim:
        linhas, _ = _find_all_in_any_frame(driver, By.XPATH, "//tr[td]", timeout=3, frame_key="diario")
        if len(linhas) >= 3:
            break
        time.sleep(0.2)

    out = []
    seen = set()
    permitidas = {"6A", "6B", "7A", "7B", "8A", "8B", "9A", "9B"}

    for tr in linhas:
        try:
            tds = tr.find_elements(By.XPATH, "./td")
            if not tds:
                continue
            txts = [((td.text or "").strip()) for td in tds]
            txts = [t for t in txts if t]
            if not txts:
                continue

            linha_txt = " | ".join(txts)
            n = normalizar_texto(linha_txt)
            if any(k in n for k in ["PESQUISAR", "TOTAL", "FILTRO", "NENHUM REGISTRO", "AÇÃO", "ACAO"]):
                continue

            sigla = _inferir_sigla_turma_por_texto(linha_txt)
            if sigla not in permitidas:
                continue

            txt_base = max(txts, key=len)
            rotulo = f"{sigla} - {txt_base}" if not txt_base.upper().startswith(sigla) else txt_base
            if rotulo not in seen:
                seen.add(rotulo)
                out.append(rotulo)
        except Exception:
            continue

    if not out:
        _salvar_diagnostico_coleta(driver, "diario_lista_vazia")

    return out

def _inferir_sigla_turma_por_texto(txt: str) -> str:
    tx = normalizar_texto(txt)

    # formato direto: 6A, 7B...
    m = re.search(r"\b([6-9][AB])\b", tx)
    if m:
        return m.group(1)

    # formato: 6 ANO A / 6º ANO A / 6o ANO A
    m2 = re.search(r"\b([6-9])\s*(?:O|º)?\s*ANO\s*([AB])\b", tx)
    if m2:
        return f"{m2.group(1)}{m2.group(2)}"

    # formato com separadores
    m3 = re.search(r"\b([6-9])\s*[-_/]\s*([AB])\b", tx)
    if m3:
        return f"{m3.group(1)}{m3.group(2)}"

    # fallback no conjunto permitido
    permitidas = ["6A", "6B", "7A", "7B", "8A", "8B", "9A", "9B"]
    for sig in permitidas:
        if re.search(rf"\b{re.escape(sig)}\b", tx):
            return sig
    return ""


def _sincronizar_uma_turma_com_fallback(driver, sigla: str, texto_base: str) -> Tuple[bool, str]:
    """Sincroniza uma turma com tentativas lineares (texto base + variações) para reduzir falhas em lote."""
    candidatos = []
    if texto_base:
        candidatos.append(texto_base)

    # Reaproveita referência oficial da lista de turmas para ampliar chance de acerto no diário
    oficial = next((t.split(" - ", 1)[1].strip() for t in lista_turmas if t.startswith(sigla + " - ")), "")
    if oficial and oficial not in candidatos:
        candidatos.append(oficial)

    # fallback mínimo por sigla (o seletor já possui score + xpath por sigla/ano)
    if sigla and sigla not in candidatos:
        candidatos.append(sigla)

    ultimo_erro = ""
    for texto in candidatos:
        try:
            _abrir_diario_classe(driver)
            _selecionar_turma_like_old(driver, texto)
            _abrir_sanfona_frequencia_like_old(driver)
            _sincronizar_alunos_da_turma_no_siepe(driver, sigla)
            return True, ""
        except Exception as exc:
            ultimo_erro = str(exc)
            continue

    return False, (ultimo_erro or f"falha ao sincronizar turma {sigla}")


def sincronizar_base_siepe(apelido: str, senha: str, atualizar_alunos: bool = False, manter_aberto: bool = False) -> Dict[str, Any]:
    """Sincroniza turmas do professor e, opcionalmente, alunos por turma (somente sob demanda)."""
    driver = None
    perfil = carregar_json(str(ARQUIVO_PERFIL_SIEPE), {"turmas": [], "updated_at": ""})
    erros: List[str] = []
    try:
        driver = _build_driver_visible()
        apelido, senha = _login_educadores_com_retry(driver, apelido, senha)
        _abrir_diario_classe(driver)

        turmas_texto: List[str] = []
        for _ in range(3):
            turmas_texto = _listar_turmas_no_diario(driver)
            if turmas_texto:
                break
            time.sleep(1)

        if not turmas_texto:
            _salvar_diagnostico_coleta(driver, "sem_turmas")
            # fallback linear/unificado: usa lista oficial local (sem abrir quadro de horários)
            turmas_texto = [t.split(" - ", 1)[1].strip() for t in lista_turmas if " - " in t]

        if not turmas_texto:
            raise RuntimeError("Não consegui coletar turmas no Diário de Classe. Verifique permissão do usuário no SIEPE.")

        turmas_detectadas = []
        permitidas = {"6A", "6B", "7A", "7B", "8A", "8B", "9A", "9B"}
        for t in turmas_texto:
            sig = _inferir_sigla_turma_por_texto(t)
            if sig not in permitidas:
                continue
            turmas_detectadas.append({"texto": t, "sigla": sig})

        # garante unicidade por sigla e ordenação pedagógica
        ordem = {"6A":1,"6B":2,"7A":3,"7B":4,"8A":5,"8B":6,"9A":7,"9B":8}
        dedup = {}
        for item in turmas_detectadas:
            dedup[item["sigla"]] = item
        turmas_detectadas = sorted(dedup.values(), key=lambda x: ordem.get(x.get("sigla","ZZ"), 99))

        perfil = {
            "turmas": turmas_detectadas,
            "total_turmas": len(turmas_detectadas),
            "updated_at": datetime.now().isoformat(),
            "alunos_updated_at": perfil.get("alunos_updated_at", ""),
            "erros": erros,
        }

        if atualizar_alunos:
            atualizadas: List[str] = []
            for i, t in enumerate(turmas_detectadas, start=1):
                texto = t.get("texto", "")
                sigla = t.get("sigla", "")
                log_info(f"Roteiro lote ({i}/{len(turmas_detectadas)}): abrir Educador → selecionar {sigla or texto} → copiar alunos")
                if not texto and not sigla:
                    continue
                # se não inferiu sigla, tenta pela lista oficial
                if not sigla and texto:
                    for lt in lista_turmas:
                        if texto[:10].upper() in lt.upper():
                            sigla = lt.split(" - ", 1)[0].strip().upper()
                            break
                if not sigla:
                    erros.append(f"sem_sigla:{texto}")
                    continue

                ok, erro = _sincronizar_uma_turma_com_fallback(driver, sigla, texto)
                if ok:
                    atualizadas.append(sigla)
                else:
                    erros.append(f"{sigla}:{erro}")

            perfil["alunos_updated_at"] = datetime.now().isoformat()
            perfil["turmas_com_alunos_atualizadas"] = atualizadas
            perfil["total_turmas_alunos_atualizadas"] = len(atualizadas)
            perfil["erros"] = erros

        salvar_json(str(ARQUIVO_PERFIL_SIEPE), perfil)
        return perfil
    finally:
        if driver:
            try:
                if manter_aberto:
                    log_warn("Sessão do SIEPE mantida aberta para inspeção. Feche o navegador e pressione ENTER para continuar.")
                    input("ENTER para finalizar sessão SIEPE...")
                driver.quit()
            except Exception:
                pass


def relatorio_semanal_professor(grade: dict, data_foco: datetime):
    """Entrega quadro semanal por turma e quantas aulas já foram dadas na semana."""
    semana_ini = data_foco
    while semana_ini.weekday() != 0:
        semana_ini = semana_ini - timedelta(days=1)
    dias_semana = [semana_ini + timedelta(days=i) for i in range(5)]

    aulas_planejadas = {}
    for d in range(5):
        for h in HORARIOS_PADRAO:
            turma = (grade.get(str(d), {}) or {}).get(h, "AULA ATIVIDADE")
            if turma == "AULA ATIVIDADE":
                continue
            aulas_planejadas[turma] = aulas_planejadas.get(turma, 0) + 1

    status = carregar_json(ARQUIVO_STATUS, status_padrao())
    aulas_dadas = {}
    for turma, datas in status.get("turmas", {}).items():
        for dt in dias_semana:
            dkey = dt.strftime("%d/%m/%Y")
            slots = datas.get(dkey, {})
            if isinstance(slots, dict):
                for _slot, rec in slots.items():
                    if isinstance(rec, dict) and rec.get("chamada_feita"):
                        aulas_dadas[turma] = aulas_dadas.get(turma, 0) + 1

    perfil = carregar_json(str(ARQUIVO_PERFIL_SIEPE), {"turmas": []})
    turmas_prof = [x.get("sigla") for x in perfil.get("turmas", []) if x.get("sigla")]
    turmas_alvo = sorted(set(turmas_prof) or set(list(aulas_planejadas.keys()) + list(aulas_dadas.keys())))

    tb = Table(title="📅 Quadro Semanal do Professor")
    tb.add_column("Turma")
    tb.add_column("Aulas na semana")
    tb.add_column("Aulas dadas")
    tb.add_column("Pendentes")
    for t in turmas_alvo:
        pl = aulas_planejadas.get(t, 0)
        dd = aulas_dadas.get(t, 0)
        tb.add_row(t, str(pl), str(dd), str(max(0, pl-dd)))
    console.print(tb)

    grade_tb = Table(title="🧭 Horário semanal (F1..F7)")
    grade_tb.add_column("Horário")
    grade_tb.add_column("SEG")
    grade_tb.add_column("TER")
    grade_tb.add_column("QUA")
    grade_tb.add_column("QUI")
    grade_tb.add_column("SEX")
    for h in HORARIOS_PADRAO:
        grade_tb.add_row(
            h,
            grade.get("0", {}).get(h, "AULA ATIVIDADE"),
            grade.get("1", {}).get(h, "AULA ATIVIDADE"),
            grade.get("2", {}).get(h, "AULA ATIVIDADE"),
            grade.get("3", {}).get(h, "AULA ATIVIDADE"),
            grade.get("4", {}).get(h, "AULA ATIVIDADE"),
        )
    console.print(grade_tb)




def _mapping_dominancia_anomala(mapeamento_manual: Dict[str, str]) -> Tuple[bool, str, int, int]:
    vals = [str(v).strip().upper() for k, v in (mapeamento_manual or {}).items() if isinstance(k, str) and str(k).startswith("id:")]
    total = len(vals)
    if total < 3:
        return False, "", 0, total
    cont: Dict[str, int] = {}
    for v in vals:
        if v in PERMITIDAS_TURMAS:
            cont[v] = cont.get(v, 0) + 1
    if not cont:
        return False, "", 0, total
    turma, qtd = sorted(cont.items(), key=lambda x: x[1], reverse=True)[0]
    return (qtd / max(1, total)) >= 0.8, turma, qtd, total

def coletar_todos_quadros_siepe(apelido: str, senha: str, usar_cache: bool = True) -> List[Dict[str, Any]]:
    if usar_cache and CACHE_QUADROS.exists():
        return json.loads(CACHE_QUADROS.read_text(encoding="utf-8"))

    driver = None
    try:
        driver = _build_driver_visible()
        apelido, senha = _login_educadores_com_retry(driver, apelido, senha)
        _abrir_quadro_horarios(driver)
        ids = _listar_ids_quadros(driver)
        if not ids:
            raise RuntimeError("Não encontrei turmas (detalheQuadroDeHorario).")

        coletados = []
        for i, (qid, label) in enumerate(ids, start=1):
            console.print(f"🔎 ({i}/{len(ids)}) Abrindo {qid} | {label}")
            _abrir_detalhe_por_id_js(driver, qid)
            titulo, grade, matriz, contexto = _extrair_quadros_e_matriz_do_detalhe(driver)
            coletados.append({"id": qid, "label": label, "titulo": titulo, "grade": grade, "matriz": matriz, "contexto_turma": contexto})
            _voltar_para_lista_turmas(driver)
            time.sleep(0.25)

        CACHE_QUADROS.write_text(json.dumps(coletados, ensure_ascii=False, indent=2), encoding="utf-8")
        return coletados
    finally:
        if driver:
            try:
                driver.quit()
            except Exception:
                pass


def importar_grade_siepe_otimizada(grade_atual: dict, apelido: str, senha: str) -> dict:
    """Importa grade do SIEPE com foco na grade do professor selecionado."""
    try:
        usar_cache = ""
        if CACHE_QUADROS.exists():
            usar_cache = input("Usar cache de horários já coletado? [S/n]: ").strip().lower()

        coletados = coletar_todos_quadros_siepe(apelido, senha, usar_cache=usar_cache != "n")
        if not coletados:
            log_warn("Sem quadros coletados. Mantenho grade atual.")
            return grade_atual

        prev = Table(title="Preview da Grade Detectada (antes de aplicar)")
        prev.add_column("ID")
        prev.add_column("Label")
        prev.add_column("Turma inferida")
        prev.add_column("Slots")
        for item in coletados[:30]:
            grade_item = item.get("grade") or {}
            total_slots = sum(len(grade_item.get(d) or {}) for d in DIAS_COL)
            resolved = _resolver_turma_item(item, mapeamento_manual if "mapeamento_manual" in locals() else {}) or inferir_turma_do_quadro(item) or "(mapear)"
            prev.add_row(str(item.get("id", "")), str(item.get("label", ""))[:44], resolved, str(total_slots))
        console.print(prev)

        professores = []
        if MEM_PROFS.exists():
            try:
                professores = json.loads(MEM_PROFS.read_text(encoding="utf-8"))
            except Exception:
                professores = []
        if not professores:
            professores = atualizar_memoria_professores(coletados)

        mapeamento_manual = carregar_json(str(MAPEAMENTO_TURMAS), {})
        if not isinstance(mapeamento_manual, dict):
            mapeamento_manual = {}

        # fluxo linear/unificado: limpar histórico de mapeamento legado para priorizar turma_meta do detalhe
        mapeamento_manual = {}
        salvar_json(str(MAPEAMENTO_TURMAS), mapeamento_manual)
        log_info("Mapeamento manual zerado (priorizando turma_meta do detalhe).")

        for item in coletados:
            qid = str(item.get("id", "")).strip()
            lbl = str(item.get("label", "")).strip()
            if not lbl and not qid:
                continue
            if qid and mapeamento_manual.get(f"id:{qid}"):
                continue
            if _resolver_turma_item(item, mapeamento_manual):
                continue
            escolha = tui_select(
                "Mapear turma não inferida",
                f"Quadro: {lbl}",
                [(k, k) for k in sorted(banco_alunos.keys())] + [("skip", "Ignorar")],
                "skip",
            )
            if escolha and escolha != "skip" and qid:
                mapeamento_manual[f"id:{qid}"] = escolha
        salvar_json(str(MAPEAMENTO_TURMAS), mapeamento_manual)

        professor = prompt("Professor (TAB): ", completer=ProfessorCompleter(professores)).strip()
        if not professor:
            return grade_atual

        periodo = tui_select("Importar grade", "Aplicar qual período da grade?", [("manha", "Manhã"), ("tarde", "Tarde"), ("tudo", "Tudo"), ("voltar", "Voltar")], "tudo")
        if periodo in {None, NAV_BACK, NAV_MENU, "voltar"}:
            return grade_atual

        nova_grade, horarios_reais, conflitos = montar_grade_professor(coletados, professor, mapeamento_manual)
        ajustar_horarios_padrao_dinamico(horarios_reais)

        grade_painel: Dict[str, Dict[str, str]] = {}
        for d in range(5):
            dd = str(d)
            grade_painel[dd] = {}
            for h in HORARIOS_PADRAO:
                if periodo == "manha" and int(h[:2]) >= 12:
                    grade_painel[dd][h] = "AULA ATIVIDADE"
                    continue
                if periodo == "tarde" and int(h[:2]) < 12:
                    grade_painel[dd][h] = "AULA ATIVIDADE"
                    continue
                grade_painel[dd][h] = (nova_grade.get(dd, {}) or {}).get(h, "AULA ATIVIDADE")

        salvar_json(ARQUIVO_HORARIOS, grade_painel)

        resumo = Table(title="Resultado da Importação por Professor")
        resumo.add_column("Dia")
        resumo.add_column("Aulas atribuídas")
        for i, dia in enumerate(DIAS_SEMANA):
            count = sum(1 for _h, turma in grade_painel[str(i)].items() if turma != "AULA ATIVIDADE")
            resumo.add_row(dia, str(count))
        console.print(resumo)

        por_turma = {}
        for d in range(5):
            for _h, turma in grade_painel[str(d)].items():
                if turma == "AULA ATIVIDADE":
                    continue
                for t in [x.strip() for x in turma.split("|") if x.strip()]:
                    por_turma[t] = por_turma.get(t, 0) + 1
        tb_turmas = Table(title="Turmas detectadas e carga semanal")
        tb_turmas.add_column("Turma")
        tb_turmas.add_column("Aulas/semana")
        for t in sorted(por_turma.keys()):
            tb_turmas.add_row(t, str(por_turma[t]))
        console.print(tb_turmas)

        if conflitos:
            console.print(Panel("\n".join(conflitos[:12]), title="⚠️ Conflitos de horário detectados", border_style="yellow"))

        log_info("Grade do professor importada e aplicada ao menu atual.")
        return grade_painel

    except Exception as exc:
        log_error(f"Erro ao importar grade do SIEPE: {exc}")
        return grade_atual


def fluxo_linear_professor(grade_atual: dict, apelido: str, senha: str) -> dict:
    """Fluxo unificado (linear): sincronizar -> importar grade -> relatório semanal."""
    console.print(Panel(
        "[bold]Fluxo linear do professor[/bold]\n"
        "1) Sincronizar turmas no SIEPE\n"
        "2) (Opcional) Sincronizar alunos por turma\n"
        "3) Importar grade por professor\n"
        "4) Exibir relatório semanal de aulas\n",
        title="🚀 Assistente Linear",
        border_style="cyan",
    ))

    manter = tui_select(
        "Fluxo Linear",
        "Deseja manter o SIEPE aberto ao final das sincronizações para depuração?",
        [("n", "Não"), ("s", "Sim")],
        "n",
    ) == "s"

    sync_alunos = tui_select(
        "Fluxo Linear",
        "Sincronizar também a lista de alunos de cada turma? (mais demorado)",
        [("n", "Não (somente turmas)"), ("s", "Sim (turmas + alunos)")],
        "n",
    ) == "s"

    try:
        perfil = sincronizar_base_siepe(apelido, senha, atualizar_alunos=sync_alunos, manter_aberto=manter)
        log_info(f"Etapa 1 concluída: {perfil.get('total_turmas', 0)} turmas detectadas.")
    except Exception as exc:
        log_error(f"Falha na sincronização inicial: {exc}")
        cont = tui_select("Fluxo Linear", "Deseja continuar mesmo assim?", [("n", "Não"), ("s", "Sim")], "n")
        if cont != "s":
            return grade_atual

    grade_nova = importar_grade_siepe_otimizada(grade_atual, apelido, senha)
    if grade_nova == grade_atual:
        log_warn("A grade permaneceu inalterada após a importação. Mantendo configuração atual.")
    else:
        aulas_ativas = 0
        for d in range(5):
            for _h, turma in (grade_nova.get(str(d), {}) or {}).items():
                if turma != "AULA ATIVIDADE":
                    aulas_ativas += 1
        log_info(f"Etapa 3 concluída: grade atualizada com {aulas_ativas} aulas ativas.")

    relatorio_semanal_professor(grade_nova, datetime.now())
    input("ENTER para voltar ao painel...")
    return grade_nova


def rotina_linear_unificada(grade_atual: dict, apelido: str, senha: str) -> dict:
    """Rotina única: usa preferências salvas (cookie) para sugerir caminho padrão por dia/horário."""
    agora = datetime.now()
    turma_sigla, slot = obter_turma_prioritaria(agora, grade_atual)
    turma_pre = next((t for t in lista_turmas if turma_sigla and t.startswith(turma_sigla + " - ")), None)

    desc_turma = turma_pre or (f"{turma_sigla} (sem rótulo completo)" if turma_sigla else "não definida")
    console.print(Panel(
        "[bold]Rotina linear unificada[/bold]\n"
        f"Dia/horário foco: {DIAS_SEMANA[min(max(agora.weekday(), 0), 4)]} | {slot or '-'}\n"
        f"Turma prioritária (cookie): {desc_turma}\n"
        "Fluxo: 1) Executar chamada da turma foco  2) (Opcional) Importar turma foco  3) (Opcional) Importar turmas do SIEPE",
        title="🧩 Rotina Unificada",
        border_style="cyan",
    ))

    acao = tui_select(
        "Rotina Unificada",
        "Escolha o fluxo padrão",
        [("seguir", "Seguir rotina padrão agora"), ("importar_turma", "Importar alunos da turma foco"), ("importar_grade", "Importar turmas do SIEPE"), ("cancel", "Voltar")],
        "seguir",
    )
    if acao in {None, NAV_BACK, NAV_MENU, "cancel"}:
        return grade_atual

    if acao == "importar_turma":
        importar_alunos_turma_via_chamada(apelido, senha, turma_pre)
        registrar_trilha_usuario(agora, "importar_turma_foco", turma_sigla, slot)
        return grade_atual

    if acao == "importar_grade":
        registrar_trilha_usuario(agora, "importar_turmas_siepe", turma_sigla, slot)
        return importar_grade_siepe_otimizada(grade_atual, apelido, senha)

    if not turma_pre:
        log_warn("Nenhuma turma foco encontrada no cookie/grade para este horário. Selecione manualmente.")
    wizard_chamada(grade_atual, agora, turma_pre, apelido, senha, slot or "manual")
    if turma_sigla:
        registrar_trilha_usuario(agora, "chamada_rotina", turma_sigla, slot)
    return grade_atual


# ==========================================================
# CHECKLIST CENTRAL + PAINEL + CHAMADA + PENDÊNCIAS
# ==========================================================
def checklist_add_tarefa(lembretes: dict, sigla: str, tarefa: TarefaChecklist):
    lembretes.setdefault("turma", {}).setdefault(sigla, []).append(asdict(tarefa))


def checklist_toggle(lembretes: dict, sigla: str, idx: int):
    tarefas = lembretes.setdefault("turma", {}).setdefault(sigla, [])
    if 0 <= idx < len(tarefas):
        tarefas[idx]["concluida"] = not tarefas[idx].get("concluida", False)


def logica_checklist_central(grade, data_foco: datetime, lembretes: dict):
    dia_idx = data_foco.weekday()
    if dia_idx > 4:
        log_warn("Sem aulas hoje. Você pode organizar checklist por turma.")
    turmas_hoje = []
    if dia_idx <= 4:
        for h in HORARIOS_PADRAO:
            t = grade.get(str(dia_idx), {}).get(h, "AULA ATIVIDADE")
            if t != "AULA ATIVIDADE" and t not in turmas_hoje:
                turmas_hoje.append(t)

    opções = [(t, t) for t in (turmas_hoje or sorted(banco_alunos.keys()))] + [("voltar", "Voltar")]
    turma = tui_select("📘 Central da Turma", "Selecione a turma", opções)
    if turma in {None, NAV_BACK, NAV_MENU, "voltar"}:
        return lembretes

    while True:
        tarefas = lembretes.setdefault("turma", {}).setdefault(turma, [])
        linhas = []
        for i, t in enumerate(tarefas):
            ic = "✅" if t.get("concluida") else "☐"
            when = f" ({t.get('data') or '-'} {t.get('horario') or ''})"
            linhas.append((str(i), f"{ic} {t.get('texto')}{when}"))
        acao = tui_select(
            f"Checklist {turma}",
            "Gerenciar tarefas",
            [("add", "Adicionar tarefa"), ("critical", "Sugerir caso crítico"), ("toggle", "Marcar/desmarcar")]
            + linhas[:12]
            + [("clear", "Limpar concluídas"), ("voltar", "Voltar")],
            "add",
        )
        if acao in {None, NAV_BACK, NAV_MENU, "voltar"}:
            return lembretes
        if acao == "add":
            texto = input("Texto da tarefa: ").strip()
            if not texto:
                continue
            data_ref = input("Data (DD/MM opcional): ").strip()
            hora_ref = input("Vincular horário (HH:MM opcional): ").strip()
            checklist_add_tarefa(lembretes, turma, TarefaChecklist(texto=texto, data=data_ref, horario=hora_ref))
        elif acao == "critical":
            checklist_add_tarefa(lembretes, turma, TarefaChecklist(texto="[CRÍTICO] Conferir aluno com faltas altas", prioridade=1))
        elif acao == "toggle":
            idx = input("Índice da tarefa: ").strip()
            if idx.isdigit():
                checklist_toggle(lembretes, turma, int(idx))
        elif acao == "clear":
            lembretes["turma"][turma] = [t for t in tarefas if not t.get("concluida")]
        elif acao.isdigit():
            checklist_toggle(lembretes, turma, int(acao))


def _criar_tabela_painel_tatico(grade, data_foco: datetime, status: dict, data_key: str) -> Table:
    """Versão compacta para ficar ao lado direito do painel de atalhos."""
    agora_idx, prox_idx = agora_slot_info(data_foco)

    tabela = Table(show_header=True, header_style="bold yellow", title="Painel Tático", box=box.SIMPLE_HEAVY)
    tabela.add_column("HORÁRIO", style="cyan", width=14, no_wrap=True)
    for d in ["SEG", "TER", "QUA", "QUI", "SEX"]:
        tabela.add_column(d, justify="center", width=7, no_wrap=True)

    linhas_destacadas = {"09:10 - 10:00", "11:10 - 12:00"}
    linha_divisoria = ["[dim]" + "─" * 12 + "[/dim]"] + ["[dim]" + "─" * 5 + "[/dim]"] * 5

    for i, h in enumerate(HORARIOS_PADRAO):
        h_ini = h.split(" - ")[0]
        slot_txt = h_ini
        if i == agora_idx:
            slot_txt = f"🟢 {h_ini}"
        elif i == prox_idx:
            slot_txt = f"⏭️ {h_ini}"

        row = [slot_txt]
        dia_atual = data_foco.weekday()
        for d in range(5):
            turma = grade.get(str(d), {}).get(h, "AULA ATIVIDADE")
            if turma == "AULA ATIVIDADE":
                row.append("[dim]ATIV.[/dim]")
            else:
                txt = f"[bold green]{turma[:6]}[/bold green]"
                if d == dia_atual:
                    st = status.get("turmas", {}).get(turma, {}).get(data_key, {}).get(h, {})
                    nivel = int(st.get("avaliacao_nivel", 0) or 0)
                    if nivel:
                        txt = _estilo_avaliacao_texto(nivel, turma[:6])
                row.append(txt)
        tabela.add_row(*row)
        if h in linhas_destacadas:
            tabela.add_row(*linha_divisoria)

    return tabela


def _tempo_aula_status(now: datetime) -> Dict[str, Any]:
    """Calcula estado do dia letivo para o relógio (aula atual/próxima/finalizado)."""
    if now.weekday() > 4:
        return {"modo": "off", "slot": "", "restante": 0, "proxima": ""}

    slots = []
    for slot in HORARIOS_PADRAO:
        mm = re.findall(r"(\d{2}):(\d{2})", slot)
        if len(mm) != 2:
            continue
        ini = now.replace(hour=int(mm[0][0]), minute=int(mm[0][1]), second=0, microsecond=0)
        fim = now.replace(hour=int(mm[1][0]), minute=int(mm[1][1]), second=0, microsecond=0)
        slots.append((slot, ini, fim))

    for slot, ini, fim in slots:
        if ini <= now <= fim:
            return {
                "modo": "running",
                "slot": slot,
                "restante": max(0, int((fim - now).total_seconds())),
                "proxima": "",
            }

    for slot, ini, _fim in slots:
        if now < ini:
            return {
                "modo": "next",
                "slot": "",
                "restante": max(0, int((ini - now).total_seconds())),
                "proxima": slot,
            }

    return {"modo": "done", "slot": "", "restante": 0, "proxima": ""}


def _flip_clock_text(dt: Optional[datetime] = None) -> str:
    """Renderiza temporizador legível para aula atual ou próxima aula."""
    now = dt or datetime.now()
    info = _tempo_aula_status(now)

    restante_seg = int(info.get("restante", 0) or 0)
    hh = restante_seg // 3600
    mm = (restante_seg % 3600) // 60
    ss = restante_seg % 60
    hhmmss = f"{hh:02d}:{mm:02d}:{ss:02d}"

    # ASCII seguro para qualquer fonte de terminal (evita blocos invisíveis)
    flip = {
        "0": (" __ ", "|  |", "|__|"),
        "1": ("    ", "   |", "   |"),
        "2": (" __ ", " __|", "|__ "),
        "3": ("__  ", " __|", "__|"),
        "4": ("    ", "|__|", "   |"),
        "5": (" __ ", "|__ ", " __|"),
        "6": (" __ ", "|__ ", "|__|"),
        "7": ("___ ", "  / ", " /  "),
        "8": (" __ ", "|__|", "|__|"),
        "9": (" __ ", "|__|", " __|"),
        ":": ("    ", " .. ", " .. "),
    }

    row1, row2, row3 = [], [], []
    for ch in hhmmss:
        a, b, c = flip.get(ch, (" ?? ", " ?? ", " ?? "))
        row1.append(a)
        row2.append(b)
        row3.append(c)

    if info["modo"] == "running":
        status_linha = f"[cyan]Aula atual: {info['slot']}[/cyan]"
        rodape = "[dim]Tempo restante para acabar a aula[/dim]"
    elif info["modo"] == "next":
        status_linha = f"[yellow]Próxima aula: {info['proxima']}[/yellow]"
        rodape = "[dim]Contagem até o início da próxima aula[/dim]"
    elif info["modo"] == "done":
        status_linha = "[green]Aulas do dia encerradas[/green]"
        rodape = "[dim]Sem aulas restantes hoje[/dim]"
    else:
        status_linha = "[dim]Sem aula em andamento[/dim]"
        rodape = "[dim]Fim de semana/recesso[/dim]"

    return "\n".join([
        f"[bold white]{hhmmss}[/bold white]",
        " ".join(row1),
        " ".join(row2),
        " ".join(row3),
        status_linha,
        rodape,
    ])


def exibir_painel_tatico(grade, data_foco: datetime, lembretes: dict):
    console.clear()
    status = carregar_json(ARQUIVO_STATUS, status_padrao())
    pend = carregar_json(ARQUIVO_PENDENCIAS, pendencias_padrao())
    pend_count = len([p for p in pend if p.get("status") == "pendente"])
    dia_idx = data_foco.weekday()
    data_key = data_foco.strftime("%d/%m/%Y")

    turmas_hoje = 0
    alertas = 0
    pend_hoje = 0
    if dia_idx <= 4:
        for h in HORARIOS_PADRAO:
            t = grade.get(str(dia_idx), {}).get(h, "AULA ATIVIDADE")
            if t != "AULA ATIVIDADE":
                turmas_hoje += 1
                if not slot_chamada_feita(status, t, data_key, h):
                    pend_hoje += 1
                total = len(banco_alunos.get(t, []))
                st = status.get("turmas", {}).get(t, {}).get(data_key, {}).get(h, {})
                falt = int(st.get("faltosos", 0))
                if total and (falt / max(total, 1) >= 0.35):
                    alertas += 1

    semana_txt = calcular_semana_atual(data_foco)
    semana_num = re.search(r"(\d+)$", semana_txt)
    sl = semana_num.group(1) if semana_num else "0"
    dia_nome = ["Seg", "Ter", "Qua", "Qui", "Sex", "Sáb", "Dom"][data_foco.weekday()]
    barra_superior = (
        f"[bold cyan]MEGATRON[/bold cyan] • "
        f"[white]{dia_nome} {data_foco.strftime('%d/%m/%Y')}[/white] • "
        f"[cyan]SL{sl}[/cyan] • "
        f"[white]{turmas_hoje} aulas[/white] • "
        f"[yellow]{pend_hoje}⏳[/yellow] • "
        f"[red]{alertas}⚠[/red] • "
        f"[magenta]{pend_count}📤[/magenta]"
    )
    console.print(f"\n{barra_superior}")

    quadro = _criar_tabela_painel_tatico(grade, data_foco, status, data_key)
    quadro.title = "Quadro Semanal"
    console.print(quadro)

    agora_idx, prox_idx = agora_slot_info(data_foco)
    agora_txt = "Agora: --"
    prox_txt = "Próxima: --"
    if dia_idx <= 4 and agora_idx is not None:
        hs = HORARIOS_PADRAO[agora_idx]
        turma = grade.get(str(dia_idx), {}).get(hs, "AULA ATIVIDADE")
        status_slot = "realizada" if slot_chamada_feita(status, turma, data_key, hs) else "pendente"
        agora_txt = f"Agora: {hs[:5]} — {turma} ({status_slot})"
    if dia_idx <= 4 and prox_idx is not None:
        hp = HORARIOS_PADRAO[prox_idx]
        turma_p = grade.get(str(dia_idx), {}).get(hp, "AULA ATIVIDADE")
        prox_txt = f"Próxima: {hp[:5]} — {turma_p}"
    console.print(f"[bold]{agora_txt}[/bold]")
    console.print(f"[dim]{prox_txt}[/dim]")

    painel_clock = Panel(_flip_clock_text(), title="⏳ FIM DA AULA", border_style="blue", width=56)
    console.print(Align.center(painel_clock))

    tarefas = []
    if dia_idx <= 4:
        vistas = []
        for h in HORARIOS_PADRAO:
            t = grade.get(str(dia_idx), {}).get(h, "AULA ATIVIDADE")
            if t not in vistas and t != "AULA ATIVIDADE":
                vistas.append(t)
                for obj in lembretes.get("turma", {}).get(t, []):
                    ic = "✅" if obj.get("concluida") else "☐"
                    tarefas.append(f"{t} {ic} {obj.get('texto')}")
    if tarefas:
        console.print(Panel(" | ".join(tarefas[:4]), title="📝 Operacional", border_style="dim"))

def _ler_resumo_frequencia_excel(turma_sigla: str, data_str: str) -> Dict[str, int]:
    """Lê a coluna da data no diário da turma e conta P/F para refletir no painel rápido."""
    try:
        arq = pasta_excel_destino() / "Faltas_Megatron.xlsx"
        if not arq.exists():
            return {"tem_registro": 0, "presentes": 0, "faltosos": 0}

        wb = openpyxl.load_workbook(arq, read_only=True, data_only=True)
        try:
            aba = f"{turma_sigla} - Diário"
            if aba not in wb.sheetnames:
                return {"tem_registro": 0, "presentes": 0, "faltosos": 0}

            ws = wb[aba]
            col_data = None
            for c in range(2, ws.max_column + 1):
                if str(ws.cell(1, c).value or "").strip() == data_str:
                    col_data = c
                    break
            if col_data is None:
                return {"tem_registro": 0, "presentes": 0, "faltosos": 0}

            presentes = 0
            faltosos = 0
            for r in range(2, ws.max_row + 1):
                v = str(ws.cell(r, col_data).value or "").strip().upper()
                if v == "P":
                    presentes += 1
                elif v == "F":
                    faltosos += 1

            tem = 1 if (presentes + faltosos) > 0 else 0
            return {"tem_registro": tem, "presentes": presentes, "faltosos": faltosos}
        finally:
            wb.close()
    except Exception:
        return {"tem_registro": 0, "presentes": 0, "faltosos": 0}


def desenhar_menu_dinamico(grade, data_foco: datetime):
    """Acessos compactos em uma única linha (layout home reorganizado)."""
    dia_idx = str(data_foco.weekday())
    grade_hoje = grade.get(dia_idx, {})
    atalhos = {f"f{i+1}": grade_hoje.get(h, "AULA ATIVIDADE") for i, h in enumerate(HORARIOS_PADRAO)}

    console.print("[bold]C[/bold] Chamada   [bold]P[/bold] Pendências   [bold]A[/bold] Anotações   [bold]G[/bold] Grade   [bold]S[/bold] Sync   [bold]R[/bold] Rotina   [bold]ESC[/bold] Sair")
    return atalhos


def escutar_atalho(teclas_validas):
    while True:
        event = keyboard.read_event()
        if event.event_type == keyboard.KEY_DOWN:
            t = event.name.lower()
            if t in teclas_validas:
                time.sleep(0.12)
                return t


def wizard_chamada(grade, data_foco: datetime, turma_pre: Optional[str], apelido: str, senha: str, slot: str):
    turma_escolhida = turma_pre
    if not turma_escolhida:
        turma_escolhida = prompt("👉 Qual turma? (TAB): ", completer=TurmaCompleter()).strip().upper()
    if not turma_escolhida or " - " not in turma_escolhida:
        log_warn("Turma inválida/cancelada.")
        return

    sigla_turma = turma_escolhida.split(" - ", 1)[0].strip()
    registrar_trilha_usuario(data_foco, "wizard_turma", sigla_turma, slot)
    texto_busca_siepe = turma_escolhida.split(" - ", 1)[1].strip()
    alunos = banco_alunos.get(sigla_turma, [])
    if not alunos:
        log_error("Turma sem alunos cadastrados.")
        return

    modo = tui_select("Wizard de Chamada", "Etapa 2/5 - Selecione modo", [("F", "Faltosos"), ("P", "Presentes"), ("U", "Falta de 1 aluno"), ("cancel", "Cancelar")], "F")
    if modo in {None, NAV_BACK, NAV_MENU, "cancel"}:
        return

    titulo = f"Etapa 3/5 - {sigla_turma} ({'Marcar FALTOSOS' if modo != 'P' else 'Marcar PRESENTES (Excel calcula FALTOSOS)'})"
    selecionados = checkbox_selector(titulo, alunos, modo)
    if selecionados == NAV_MENU:
        return
    if selecionados == NAV_BACK:
        return wizard_chamada(grade, data_foco, turma_escolhida, apelido, senha, slot)
    if selecionados is None:
        return
    selecionados = list(dict.fromkeys(selecionados))
    if modo == "U" and len(selecionados) != 1:
        log_warn("Modo U exige exatamente 1 aluno selecionado.")
        return

    if modo == "P":
        selecionados_norm = {normalizar_texto(n) for n in selecionados}
        faltosos_excel = [a for a in alunos if normalizar_texto(a) not in selecionados_norm]
    else:
        faltosos_excel = selecionados

    total = len(alunos)
    falt = len(faltosos_excel)
    resumo = (
        f"Etapa 4/5 - Pré-visualização\n"
        f"Turma: {sigla_turma}\nData: {data_foco.strftime('%d/%m/%Y')}\nModo: {modo}\n"
        f"Total alunos: {total}\nSelecionados: {len(selecionados)}\nFaltosos calculados: {falt}\n"
    )
    if falt == 0:
        resumo += "\n⚠️ ALERTA: 0 faltosos."
    if total and falt / total > 0.35:
        resumo += "\n🚨 ALERTA FORTE: >35% faltosos."

    conf = tui_select("Confirmação", resumo, [("ok", "Confirmar e executar"), ("back", "Voltar"), ("pend", "Salvar pendência"), ("cancel", "Cancelar")], "ok")
    if conf in {None, NAV_BACK, NAV_MENU, "cancel"}:
        return
    if conf in {"back", NAV_BACK}:
        return wizard_chamada(grade, data_foco, turma_escolhida, apelido, senha, slot)
    if conf == "pend":
        adicionar_pendencia(PendenciaRecord(
            turma=sigla_turma,
            data=data_foco.strftime('%d/%m/%Y'),
            modo=modo,
            lista_digitada=selecionados,
            faltosos_excel=faltosos_excel,
            observacoes="Rascunho manual",
            timestamp=datetime.now().isoformat(),
        ))
        salvar_no_excel(sigla_turma, faltosos_excel, {}, data_foco.strftime('%d/%m/%Y'), dia_letivo=True)
        log_info("Pendência salva com sucesso.")
        return

    notas_do_dia = {}
    sim_fail = tui_select("Execução", "Simular falha SIEPE para teste offline?", [("n", "Não"), ("s", "Sim")], "n") == "s"
    while True:
        try:
            dia_letivo = executar_siepe(apelido, senha, texto_busca_siepe, modo, selecionados, data_foco, force_fail=sim_fail, sigla_turma=sigla_turma)
            if not dia_letivo:
                log_warn("Dia não letivo no SIEPE. Salvo em Excel e fila de pendências.")
                adicionar_pendencia(PendenciaRecord(
                    turma=sigla_turma,
                    data=data_foco.strftime('%d/%m/%Y'),
                    modo=modo,
                    lista_digitada=selecionados,
                    faltosos_excel=faltosos_excel,
                    observacoes="Dia não letivo/retorno manual necessário",
                    timestamp=datetime.now().isoformat(),
                ))
            salvar_no_excel(sigla_turma, faltosos_excel, notas_do_dia, data_foco.strftime("%d/%m/%Y"), dia_letivo)
            status = carregar_json(ARQUIVO_STATUS, status_padrao())
            presentes = max(0, len(alunos) - len(faltosos_excel))
            media = round((presentes / len(alunos)) * 10, 2) if alunos else 0

            esc_av = {
                "1": "Ruim",
                "2": "Regular",
                "3": "Bom",
                "4": "Muito Bom",
                "5": "Excelente",
            }
            av = tui_select(
                "Avaliação da turma",
                f"Etapa final: avalie o comportamento da turma {sigla_turma}",
                [("1", "1 ⭐ - Ruim"), ("2", "2 ⭐ - Regular"), ("3", "3 ⭐ - Bom"), ("4", "4 ⭐ - Muito Bom"), ("5", "5 ⭐ - Excelente")],
                "3",
            )
            if av in {None, NAV_BACK, NAV_MENU}:
                av = "3"

            registrar_status_chamada(
                status,
                sigla_turma,
                data_foco.strftime("%d/%m/%Y"),
                slot,
                presentes,
                len(faltosos_excel),
                media,
                avaliacao_nivel=int(av),
                avaliacao_label=esc_av.get(str(av), "Bom"),
            )
            salvar_json(ARQUIVO_STATUS, status)
            log_info("Chamada concluída.")
            return
        except RuntimeError as exc:
            acao = modal_erro_siepe(str(exc))
            if acao == "retry":
                continue
            if acao == "pend":
                adicionar_pendencia(PendenciaRecord(
                    turma=sigla_turma,
                    data=data_foco.strftime('%d/%m/%Y'),
                    modo=modo,
                    lista_digitada=selecionados,
                    faltosos_excel=faltosos_excel,
                    observacoes=str(exc),
                    timestamp=datetime.now().isoformat(),
                    status="falhou",
                    erro=str(exc),
                ))
                salvar_no_excel(sigla_turma, faltosos_excel, notas_do_dia, data_foco.strftime("%d/%m/%Y"), True)
            return


def tela_pendencias(apelido: str, senha: str):
    while True:
        pend = carregar_json(ARQUIVO_PENDENCIAS, pendencias_padrao())
        tabela = Table(title="📤 Pendências de Reenvio")
        tabela.add_column("#")
        tabela.add_column("Turma")
        tabela.add_column("Data")
        tabela.add_column("Modo")
        tabela.add_column("Status")
        for i, p in enumerate(pend):
            tabela.add_row(str(i), p.get("turma", ""), p.get("data", ""), p.get("modo", ""), p.get("status", ""))
        console.print(tabela)

        ac = tui_select("Pendências", "Escolha uma ação", [("one", "Reenviar uma"), ("all", "Reenviar todas pendentes"), ("view", "Ver detalhes"), ("ok", "Marcar resolvida manualmente"), ("voltar", "Voltar")], "one")
        if ac in {None, NAV_BACK, NAV_MENU, "voltar"}:
            return
        if ac in {"one", "view", "ok"}:
            idx = input("Índice: ").strip()
            if not idx.isdigit() or int(idx) >= len(pend):
                continue
            i = int(idx)
            reg = pend[i]
            if ac == "view":
                console.print(Panel(json.dumps(reg, ensure_ascii=False, indent=2), title=f"Pendência {i}"))
                continue
            if ac == "ok":
                atualizar_pendencia(i, {"status": "enviado", "erro": "resolvido manual"})
                continue
            try:
                texto_busca = next((t.split(" - ", 1)[1] for t in lista_turmas if t.startswith(reg["turma"] + " - ")), reg["turma"])
                executar_siepe(apelido, senha, texto_busca, reg["modo"], reg["lista_digitada"], datetime.strptime(reg["data"], "%d/%m/%Y"), sigla_turma=reg.get("turma", ""))
                atualizar_pendencia(i, {"status": "enviado", "erro": ""})
            except Exception as exc:
                atualizar_pendencia(i, {"status": "falhou", "erro": str(exc)})
        if ac == "all":
            for i, reg in enumerate(pend):
                if reg.get("status") not in {"pendente", "falhou"}:
                    continue
                try:
                    texto_busca = next((t.split(" - ", 1)[1] for t in lista_turmas if t.startswith(reg["turma"] + " - ")), reg["turma"])
                    executar_siepe(apelido, senha, texto_busca, reg["modo"], reg["lista_digitada"], datetime.strptime(reg["data"], "%d/%m/%Y"), sigla_turma=reg.get("turma", ""))
                    atualizar_pendencia(i, {"status": "enviado", "erro": ""})
                except Exception as exc:
                    atualizar_pendencia(i, {"status": "falhou", "erro": str(exc)})


def importar_alunos_turma_via_chamada(apelido: str, senha: str, turma_pre: Optional[str] = None):
    """Importa alunos direto do ambiente de chamada (frequência), armazenando por turma."""
    turma_escolhida = turma_pre
    if not turma_escolhida:
        turma_escolhida = prompt("👉 Turma para importar alunos (TAB): ", completer=TurmaCompleter()).strip().upper()
    if not turma_escolhida or " - " not in turma_escolhida:
        log_warn("Turma inválida/cancelada.")
        return

    sigla_turma = turma_escolhida.split(" - ", 1)[0].strip().upper()
    texto_busca_siepe = turma_escolhida.split(" - ", 1)[1].strip()

    acao_lista = tui_select(
        "Importar Alunos",
        f"Turma {sigla_turma}: atualizar lista de alunos do SIEPE ou manter lista atual?",
        [("atualizar", "Atualizar lista de alunos"), ("manter", "Manter lista de alunos"), ("cancel", "Cancelar")],
        "atualizar",
    )
    if acao_lista in {None, NAV_BACK, NAV_MENU, "cancel"}:
        return
    if acao_lista == "manter":
        log_info(f"Lista atual mantida para {sigla_turma} ({len(banco_alunos.get(sigla_turma, []))} alunos).")
        return

    driver = None
    try:
        driver = _build_driver_visible()
        apelido, senha = _login_educadores_com_retry(driver, apelido, senha)
        ok = False
        ultimo_erro = None
        for tentativa in range(1, 4):
            try:
                _abrir_diario_classe(driver)
                _selecionar_turma_like_old(driver, texto_busca_siepe)
                _abrir_sanfona_frequencia_like_old(driver)
                _sincronizar_alunos_da_turma_no_siepe(driver, sigla_turma)
                ok = True
                break
            except StaleElementReferenceException as exc:
                ultimo_erro = exc
                time.sleep(0.4)
                continue
            except Exception as exc:
                ultimo_erro = exc
                if "stale" in normalizar_texto(str(exc)) and tentativa < 3:
                    time.sleep(0.4)
                    continue
                raise
        if not ok and ultimo_erro:
            raise RuntimeError(str(ultimo_erro))
        log_info(f"Importação linear concluída para {sigla_turma}.")
    except Exception as exc:
        log_error(f"Falha ao importar alunos da turma {sigla_turma}: {exc}")
    finally:
        if driver:
            try:
                driver.quit()
            except Exception:
                pass


def importar_todos_alunos_via_siepe(apelido: str, senha: str):
    """Fluxo linear/unificado para atualizar alunos de todas as turmas em uma única execução."""
    try:
        manter = tui_select("Depuração SIEPE", "Deseja manter o SIEPE aberto após a sincronização?", [("n", "Não"), ("s", "Sim")], "n") == "s"
        perfil = sincronizar_base_siepe(apelido, senha, atualizar_alunos=True, manter_aberto=manter)
        total = perfil.get("total_turmas", 0)
        atualizadas = int(perfil.get("total_turmas_alunos_atualizadas", len(perfil.get("turmas_com_alunos_atualizadas", []) or [])) or 0)
        log_info(f"Sincronização completa finalizada: {atualizadas}/{total} turmas com alunos atualizados.")
    except Exception as exc:
        log_error(f"Falha ao importar alunos de todas as turmas: {exc}")


def logica_anotar_avulso(data_string: str):
    turma = prompt("👉 Turma p/ anotação (TAB): ", completer=TurmaCompleter()).strip().upper()
    if not turma or " - " not in turma:
        return
    sigla = turma.split(" - ", 1)[0]
    alunos = banco_alunos.get(sigla, [])
    sels = checkbox_selector(f"Anotação em lote - {sigla}", alunos, "F")
    if sels is None or sels == NAV_BACK or sels == NAV_MENU:
        return

    escala = {
        "1": "Ruim",
        "2": "Regular",
        "3": "Bom",
        "4": "Muito Bom",
        "5": "Excelente",
    }
    nivel = tui_select(
        "Classificação do grupo",
        "Selecione o nível de avaliação",
        [
            ("1", "1 ⭐ - Ruim"),
            ("2", "2 ⭐ - Regular"),
            ("3", "3 ⭐ - Bom"),
            ("4", "4 ⭐ - Muito Bom"),
            ("5", "5 ⭐ - Excelente"),
            ("cancel", "Cancelar"),
        ],
        "3",
    )
    if nivel in {None, NAV_BACK, NAV_MENU, "cancel"}:
        return

    txt = input("Texto da anotação (opcional): ").strip()
    base = f"Avaliação {nivel} - {escala[nivel]}"
    texto_final = f"{base} | {txt}" if txt else base
    notas = {n: {"texto": texto_final, "nivel": int(nivel)} for n in sels}
    salvar_no_excel(sigla, [], notas, data_string, True)


def configurar_dia(nome_dia: str):
    console.print(f"\n[bold cyan]Configurar {nome_dia}[/bold cyan]")
    grade_dia = {}
    for h in HORARIOS_PADRAO:
        ent = prompt(f"{h} > ", completer=TurmaCompleter()).strip().upper()
        if not ent:
            grade_dia[h] = "AULA ATIVIDADE"
        elif " - " in ent:
            grade_dia[h] = ent.split(" - ", 1)[0]
        else:
            grade_dia[h] = ent
    return grade_dia


def editar_horarios(grade_atual: dict, apelido: str, senha: str):
    esc = tui_select("Configurar Horários", "Escolha", [("1", "Corrigir horário manualmente"), ("2", "Refazer semana"), ("3", "Importar turmas do SIEPE"), ("5", "Relatório semanal do professor"), ("V", "Voltar")], "3")
    if esc in {None, NAV_BACK, NAV_MENU, "V"}:
        return grade_atual
    if esc == "2":
        grade = {str(i): configurar_dia(d) for i, d in enumerate(DIAS_SEMANA)}
        salvar_json(ARQUIVO_HORARIOS, grade)
        return grade
    if esc == "1":
        d = tui_select("Dia", "Escolha dia", [(str(i), DIAS_SEMANA[i]) for i in range(5)], "0")
        if d in {None, NAV_BACK, NAV_MENU}:
            return grade_atual
        grade_atual[d] = configurar_dia(DIAS_SEMANA[int(d)])
        salvar_json(ARQUIVO_HORARIOS, grade_atual)
        return grade_atual
    if esc == "3":
        return importar_grade_siepe_otimizada(grade_atual, apelido, senha)
    if esc == "5":
        relatorio_semanal_professor(grade_atual, datetime.now())
        input("ENTER para voltar...")
        return grade_atual
    return grade_atual


def validar_ambiente_execucao() -> None:
    """Checklist linear de ambiente para reduzir falhas em execução real."""
    avisos = []
    if not Path(CRED_FILE).exists():
        avisos.append("Credenciais ainda não salvas (meu_acesso.txt será criado na primeira execução).")
    if not shutil.which("chrome") and not shutil.which("chrome.exe"):
        avisos.append("Chrome não encontrado no PATH. Verifique instalação do navegador.")
    if avisos:
        console.print(Panel("\n".join([f"• {a}" for a in avisos]), title="⚠️ Checklist de inicialização", border_style="yellow"))


# ==========================================================
# MAIN
# ==========================================================
def main():
    grade = carregar_json(ARQUIVO_HORARIOS, {str(i): {h: "AULA ATIVIDADE" for h in HORARIOS_PADRAO} for i in range(5)})
    lembretes = carregar_json(ARQUIVO_LEMBRETES, lembretes_padrao())
    data_foco = datetime.now()

    validar_ambiente_execucao()
    carregar_banco_alunos_completo()

    apelido, senha = carregar_credenciais()
    if not apelido or not senha:
        apelido, senha = pedir_credenciais_interativas()

    ultima_turma = None
    while True:
        exibir_painel_tatico(grade, data_foco, lembretes)
        atalhos = desenhar_menu_dinamico(grade, data_foco)
        teclas = ["c", "p", "a", "g", "s", "r", "esc", "f8", "f9", "f10", "f11", "f12"] + list(atalhos.keys())
        tecla = escutar_atalho(teclas)

        if tecla == "esc":
            break
        if tecla in {"g", "f9"}:
            grade = editar_horarios(grade, apelido, senha)
            continue
        if tecla in {"a", "f8"}:
            logica_anotar_avulso(data_foco.strftime("%d/%m/%Y"))
            continue
        if tecla == "f11":
            ent = input("Data (H, DD, DDMM, DD/MM, DD/MM/AAAA): ").strip()
            try:
                data_foco = parse_data_pratica(ent, data_foco)
            except Exception:
                log_error("Formato inválido")
            continue
        if tecla == "f12":
            lembretes = logica_checklist_central(grade, data_foco, lembretes)
            salvar_json(ARQUIVO_LEMBRETES, lembretes)
            continue
        if tecla == "p":
            tela_pendencias(apelido, senha)
            continue
        if tecla == "r":
            grade = rotina_linear_unificada(grade, apelido, senha)
            continue
        if tecla == "s":
            importar_todos_alunos_via_siepe(apelido, senha)
            continue

        if data_foco.weekday() > 4:
            ac = tui_select("Dia não letivo", "Hoje não é dia letivo. Escolha uma ação útil:", [("edit", "Editar horários"), ("pend", "Ver pendências"), ("rel", "Exportar relatório Excel"), ("voltar", "Voltar")], "pend")
            if ac == "edit":
                grade = editar_horarios(grade, apelido, senha)
            elif ac == "pend":
                tela_pendencias(apelido, senha)
            elif ac == "rel":
                log_info("Relatório já é atualizado no Excel com aba Resumo Semanal.")
            continue

        turma_pre = None
        slot = "manual"
        if tecla in atalhos:
            slot = HORARIOS_PADRAO[int(tecla[1:]) - 1]
            sigla = atalhos[tecla]
            if sigla == "AULA ATIVIDADE":
                log_warn("Aula atividade: sem chamada.")
                continue
            turma_pre = next((t for t in lista_turmas if t.startswith(sigla + " - ")), None)
            registrar_trilha_usuario(data_foco, "atalho_turma", sigla, slot)
        elif tecla in {"f10", "c"}:
            escolha_f10 = tui_select(
                "F10 - Ações da Turma",
                "Escolha uma ação",
                [("importar", "📥 Importar alunos da turma"), ("importar_todos", "📦 Puxar todos os alunos (todas as turmas)"), ("chamada", "🔍 Lançamento manual (chamada)"), ("voltar", "Voltar")],
                "importar",
            )
            if escolha_f10 in {None, NAV_BACK, NAV_MENU, "voltar"}:
                continue
            if escolha_f10 == "importar":
                turma_base = ultima_turma
                importar_alunos_turma_via_chamada(apelido, senha, turma_base)
                if turma_base and " - " in turma_base:
                    registrar_trilha_usuario(data_foco, "importar_alunos_turma", turma_base.split(" - ", 1)[0], "manual")
                continue
            if escolha_f10 == "importar_todos":
                importar_todos_alunos_via_siepe(apelido, senha)
                continue
            if ultima_turma:
                rep = tui_select("Manual", f"Repetir última chamada ({ultima_turma})?", [("sim", "Sim"), ("nao", "Não")], "sim")
                if rep == "sim":
                    turma_pre = ultima_turma

        wizard_chamada(grade, data_foco, turma_pre, apelido, senha, slot)
        if turma_pre:
            ultima_turma = turma_pre


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        console.print("\n🛑 Megatron encerrado.")
    except Exception as exc:
        console.print(f"\n[bold red]❌ Falha inesperada: {exc}[/bold red]")
        console.print("[yellow]Dica:[/yellow] verifique ChromeDriver, internet, credenciais e arquivos JSON.")
