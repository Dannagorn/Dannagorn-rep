import getpass
import json
import os
import time
import unicodedata
from datetime import datetime
from pathlib import Path

import pandas as pd
import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill
from prompt_toolkit import prompt
from prompt_toolkit.completion import Completer, Completion

try:
    from selenium import webdriver
    from selenium.common.exceptions import TimeoutException, WebDriverException
    from selenium.webdriver.common.by import By
    from selenium.webdriver.common.keys import Keys
    from selenium.webdriver.support import expected_conditions as EC
    from selenium.webdriver.support.ui import WebDriverWait
except Exception:  # selenium pode não estar disponível em todos os ambientes
    webdriver = None
    TimeoutException = WebDriverException = Exception
    By = Keys = EC = WebDriverWait = None

try:
    import win32com.client
except Exception:  # pywin32 só existe no Windows
    win32com = None


# ==========================================================
# ⚙️ BANCO DE DADOS (TURMAS E ALUNOS)
# ==========================================================
lista_turmas = [
    "6A - ENSINO FUNDAMENTAL DE ANOS FINAIS DE 3500H - 6º ANO A",
    "7A - ENSINO FUNDAMENTAL DE ANOS FINAIS DE 3500H - 7º ANO A",
    "8A - ENSINO FUNDAMENTAL DE ANOS FINAIS DE 3500H - 8º ANO A",
    "9A - ENSINO FUNDAMENTAL DE ANOS FINAIS DE 3500H - 9º ANO A",
    "9B - ENSINO FUNDAMENTAL DE ANOS FINAIS DE 3500H - 9º ANO B",
    "AULA ATIVIDADE",
]

banco_alunos = {
    "6A": ["ALEFF GABRIEL GOMES DA SILVA", "ALISSON TORRES DA SILVA", "ALYCIA HADASSA RAINELLE SIQUEIRA DE OLIVEIRA", "ARTHUR VINICIUS DE SOUZA MACHADO", "ASHLEY VICTORIA ANDRADE DE LUCENA", "DAVI LUCAS ROCHA DE OLIVEIRA", "ELOHA BEATRIZ SANTOS DA SILVA", "EMILLY VITORIA DAMASIO SOUZA", "ESTER ALVES DA SILVA CORDOVILLE", "EVERTON LUCAS OLIVEIRA DA SILVA", "IAGO FERNANDO ALVES DOS SANTOS", "ISADORA MARIA LINS DA PAIXAO VIEIRA", "JESSICA MARIA PEREIRA DOS SANTOS", "KAYO SAMUEL DE LIRA CAMILLO", "KEVILYN FERNANDA DE SOUZA LIMA", "LUAN MIQUEIAS JUSTINO DE MELO VIEIRA", "LUCAS VITOR FARIAS DA SILVA", "LUCIANA MARIA DOS SANTOS BARRETO", "LUCIANA VICTORYA ACIOLI CORREIA DE OLIVEIRA", "MARIA JULIA FLORENCIO DE SOUZA", "NICOLAS HENRIQUE MEDEIROS DE ALBUQUERQUE", "PEDRO FABIANO FRANCA DE SANTANA", "RAYLAN LUCAS JOSE DA SILVA", "RHANYELLY VITORIA LIMA DE ARAUJO", "SERGIO KAYKY DOS SANTOS SILVA", "SOPHIA GODOI N DA SILVA"],
    "7A": ["ADRIANA KETHYLLE MELO PERCILIO", "ALERRANDRO FERREIRA FIDELIS DA SILVA", "ANA JULIA MELO DA SILVA", "ANA KAROLINA FERREIRA DE ARAUJO", "ANTONY VICTOR LIMA DE OLIVEIRA", "ARTHUR NUNES BEZERRA", "CAIO ERNANDE BORGES PESSOA", "CÁSSIO ALEXANDRINO DA SILVA", "DAVI LUIZ SOUZA DE LIMA", "ESTEVÃO VITOR RODRIGUES", "GLEICY BEATRIZ XAVIER DE SOUZA", "GUYBSON RAFAEL LIMA DE VASCONCELOS", "ISMAELY CABRAL SALES", "JOAO RAPHAEL COSTA DE SOUZA", "JOSÉ CARLOS DOMINGOS DA SILVA", "JUAN RIQUELME FOSTER SANTOS CHAVES", "KAUÃ ALVES DANTAS", "KAUANE CONCEICAO DANTAS DE ARAUJO", "LARA GOMES DE LIMA", "LUCAS MATHEUS FERREIRA DE SANTANA", "LYARA RICHELLY FERREIRA CARDOSO", "MARIA CLARA ALVES AROUCHA", "MARIA CLARA MEDEIROS DE ARAÚJO", "MARIA GABRIELLY BATISTA MOURA", "MICAEL MAURICIO LUCIANO MENDES DA SILVA", "MIGUEL ALEXANDRE RODRIGUES DE LIMA", "MURILO HENRIQUE NUNES DA SILVA", "PAULA KEMILLY DE LIMA SANTANA", "RAISSA WILIANE GOMES DA SILVA", "RAYZA VITORIA MARTINS DO NASCIMENTO", "RHAKELLY VITORIA DE MOURA ALVES DA SILVA", "RIAN KAIO VICENTE DA SILVA", "RICHARLLYSON RODRIGUES EHRHARDT DE FRAGA", "SAMUEL GONCALVES DA SILVA", "THAUANY LETICIA BARRETO DE MORAIS"],
    "8A": ["ANA CLARA DOS SANTOS RAMOS", "ANA WYCTORYA DE ALMEIDA SOARES", "BRENDA NICOLE DOS SANTOS RAGO", "CICERO TINO DA SILVA NETO", "DAVI DOUGLAS FERREIRA DA SILVA", "ELLEN CAROLYNA RAMOS MIGUEL", "FELIPE GABRIEL PEREIRA BARRETO", "GILMAR FELIX DE SOUZA JUNIOR", "ISABELLA LARISSA DOS SANTOS SILVA", "JOÃO PETHERSON SANTANA DA SILVA", "JONATHAN VINICIUS MORAIS ALEXANDRINO", "JOSE MARCOS MATIAS DE SOUZA", "JULIA FERNANDA DA SILVA LIMA", "JULIA VITORIA LIMA DA SILVA", "KAUA VINICIUS SANTOS DA SILVA", "KAUÊ SAMUEL NASCIMENTO DE SANTANA", "KLEBER RAFAEL LIRA DE OLIVEIRA", "LAILA VITÓRIA MAFRA BEZERRA DE ALENCAR", "LUANA KASTINGN FERRAZ", "LUCAS RUAN AGUIAR DOS SANTOS LIMA", "MARIA EDUARDA DA COSTA SILVA", "MARIA JULIA BATISTA DOS SANTOS", "MARIA VALENTINA ELIAS ALEXANDRINO", "MARIA VITORIA DOS SANTOS JACINTO", "MELLYSSA ALVES DA SILVA MATA", "MILLENA VITORIA CABRAL DA SILVA", "REBECA CRISTINA BORGES LIMA FERREIRA", "RICHARD KAUA ARAUJO DA SILVA", "RIKELME FIDELIS DA SILVA", "RYAN KALLEB GERMANO ARAUJO", "SOPHIA GONÇALVES DE OLIVEIRA", "TALERRANDRO WELINILTON TAVARES DE FREITAS", "THAYS THAMYRES MORAES GADELHA", "WALISON FELIPE CARNEIRO DA SILVA", "ARTHUR HENRIQUE AYRES BARRETO"],
    "9A": ["ALEXSANDRO DOS SANTOS DE LIMA", "ALLANA VICTORIA SANTOS DE OLIVEIRA", "ANA BEATRIZ DA SILVA LEANDRO", "ANA BEATRIZ DA SILVA PEREIRA", "ANA CLARA DA SILVA", "ANA LUIZA DA SILVA PEREIRA", "ANA REBEKA DOS ANJOS MENDES DA SILVA", "BIANCA VICTORIA SILVA SOUZA NUNES", "CLEIBER ANTONIO BEZERRA DA SILVA JUNIOR", "DAVI EMANUEL DA SILVA", "EMANUEL FERNANDES DA COSTA SILVA", "EMANUELA QUEZIA SANTANA DE MORAES", "EVELLY RAYANNE DE ALMEIDA FERREIRA", "FABIO GABRIEL MOURA DA SILVA", "HELOÍSE GABRIELLE MARQUES BARRETO", "JOÃO GUILHERME VITAL BEZERRA", "JONATAN GABRIEL NASCIMENTO DOS SANTOS", "JONATHAN MIGUEL CORREIA DE SOUZA", "JOSICLAY KAUA DE CASTRO SILVA", "JULIA YASMIM NASCIMENTO DA SILVA", "JULLIA VITORIA PEREIRA NASCIMENTO", "JULLYA GHANDY SANTOS EISENHOWER MATOS", "KAIK DE OLIVEIRA MENDONÇA", "KAUANY FERNANDA FARIAS DA SILVA", "KLEBISON HENRIQUE OLIVEIRA DO NASCIMENTO", "LARA HUANE DOS ANJOS", "LUCAS GOMES DE LIMA BRITO", "LUIZ ARTHUR COSTA", "LUIZA MIRELLA MATIAS DA SILVA", "MARIA ELLEN GOMES DA SILVA", "NATHALIA DE ALMEIDA DE ARAUJO MOURA", "RAIZA KAUANNA AQUINO PASSOS DA SILVA", "RHAFAELA VITORIA DA SILVA GOMES", "RUTHIELLY GOMES DE SANTANA", "SAMARA VITORIA SOARES SILVA", "SAULO CHAVES DO NASCIMENTO JUNIOR", "TARCILA ALMEIDA DE VASCONCELOS", "THAIANNY STEPHANIE RIBEIRO PESSOA DE MORAIS", "THALLYTA EVELYN DOS SANTOS", "VINICIOS MATEUS ALVES AROUCHA", "WILLAMS HENRIQUE BATISTA DA COSTA", "WILLYANE NATASHA LIMA DO CARMO"],
    "9B": ["ALEXSANDRO VIDAL DA SILVA", "ALISSON ALVES DE LIMA", "ANA ROBERTA MOURA MENDES DA SILVA", "ATHOS WINICIUS JOSE ALVES DA CUNHA", "BRUNIELLY KIMBERLY FRANCA DA SILVA SANTOS", "CARLOS ADEMIR GOMES DE OLIVEIRA", "DIOGO RAPHAEL GOMES DA CUNHA", "FÁBIO MIGUEL ROQUE SIMÕES", "GESSICA MAYARA GOMES DA SILVA", "INGRIDI VITORIA FERREIRA DE MELO", "JAYLANNY MYCHELI DA SILVA ARAUJO", "JEISSON GABRIEL ALVES RODRIGUES", "JENIFFER VITORIA SALES DA SILVA", "JOSÉ GUSTAVO DA SILVA DE VASCONCELOS", "JÚLIA GABRIELLY DA CRUZ ALVES", "KAIC JOSÉ GOMES DA SILVA", "KAIKI RIAN GOMES DE SANTANA", "LARISSA BIANCA SOARES DA SILVA", "LUCAS GABRIEL GOMES CAVALCANTI", "LUCCAS DANIEL CATUNDA XAVIER", "LUIZ FERNANDO DE LIMA SANTOS", "MARIA APARECIDA MENEZES BATISTA", "MARIA CECÍLIA OLIVEIRA RIBEIRO", "MARIA GABRIELLA CORREIA LIMA", "MARIA VITORIA CORREIA NASCIMENTO", "MARÍLIA VITÓRIA DOS SANTOS GOMES", "MIGUEL LUCAS DA SILVA MELO", "MIKAELLY GOMES DA SILVA", "NAYANNA MARQUES DA SILVA", "RAFAEL OLIVEIRA DA SILVA", "RICHARLYSON DOUGLAS DOS SANTOS SOARES", "SHEYLLA NATALI MEDEIROS DA SILVA", "THAYZA NAYARA VIEIRA DA SILVA", "WEMILY RAUANY RIBEIRO DA SILVA"],
}

ARQUIVO_HORARIOS = "horarios_megatron.json"
HORARIOS_PADRAO = ["07:30 - 08:20", "08:20 - 09:10", "09:10 - 10:00", "10:20 - 11:10", "11:10 - 12:00", "13:00 - 13:50", "13:50 - 14:40"]
DIAS_SEMANA = ["Segunda", "Terça", "Quarta", "Quinta", "Sexta"]
CRED_FILE = Path(".megatron_acesso.json")


def normalizar_texto(texto: str) -> str:
    return unicodedata.normalize("NFKD", texto.upper()).encode("ASCII", "ignore").decode("utf-8").strip()


class TurmaCompleter(Completer):
    def get_completions(self, document, complete_event):
        texto = document.text.upper()
        for turma in lista_turmas:
            if texto in turma:
                yield Completion(turma, start_position=-len(document.text))


class TurmasMultiCompleter(Completer):
    def get_completions(self, document, complete_event):
        partes = document.text.upper().split(",")
        trecho = partes[-1].lstrip()
        if not trecho:
            return
        for turma in lista_turmas:
            if trecho in turma:
                yield Completion(turma + ", ", start_position=-len(trecho))


class AlunoCompleter(Completer):
    def __init__(self, lista_alunos):
        self.alunos = lista_alunos

    def get_completions(self, document, complete_event):
        texto = document.text.upper()
        for aluno in self.alunos:
            if texto in aluno:
                yield Completion(aluno, start_position=-len(document.text))


class MegatronCompleter(Completer):
    def __init__(self, lista_alunos):
        self.alunos = lista_alunos

    def get_completions(self, document, complete_event):
        partes = document.text.upper().split(",")
        trecho_atual = partes[-1].lstrip()
        concluidos = {p.strip() for p in partes[:-1] if p.strip()}
        if not trecho_atual:
            return
        for aluno in self.alunos:
            if aluno not in concluidos and trecho_atual in aluno:
                yield Completion(aluno + ", ", start_position=-len(trecho_atual))


def carregar_horarios():
    if not os.path.exists(ARQUIVO_HORARIOS):
        return {}
    try:
        with open(ARQUIVO_HORARIOS, "r", encoding="utf-8") as arquivo:
            dados = json.load(arquivo)
            return dados if isinstance(dados, dict) else {}
    except (json.JSONDecodeError, OSError):
        return {}


def salvar_horarios(grade):
    with open(ARQUIVO_HORARIOS, "w", encoding="utf-8") as arquivo:
        json.dump(grade, arquivo, ensure_ascii=False, indent=4)


def exibir_tabela_horarios(grade):
    print("\n" + "📅 SUA GRADE SEMANAL ATUAL ".center(100, "="))
    header = f"| {'HORÁRIO':<13} | {'SEG':<14} | {'TER':<14} | {'QUA':<14} | {'QUI':<14} | {'SEX':<14} |"
    print("-" * len(header))
    print(header)
    print("-" * len(header))
    for horario in HORARIOS_PADRAO:
        linha = f"| {horario:<13} "
        for i in range(5):
            turma = grade.get(str(i), {}).get(horario, "AULA ATIVIDADE")
            linha += f"| {turma[:14]:<14} "
        linha += "|"
        print(linha)
    print("-" * len(header))


def configurar_dia(nome_dia):
    grade_dia = {}
    entrada = prompt(f"  {nome_dia.upper()} > ", completer=TurmasMultiCompleter()).strip()
    partes = [p.strip() for p in entrada.split(",")]
    for idx, horario in enumerate(HORARIOS_PADRAO):
        if idx < len(partes) and partes[idx]:
            grade_dia[horario] = partes[idx].split(" - ", 1)[0].strip().upper()
        else:
            grade_dia[horario] = "AULA ATIVIDADE"
    return grade_dia


def configurar_horarios():
    print("\n" + "⚙️ CONFIGURAÇÃO COMPLETA DA GRADE ".center(65, "="))
    print("Digite as turmas em sequência, separadas por VÍRGULA. (Use TAB para completar).")
    print("Exemplo: 9A, 9B, , 7A")
    grade = {str(i): configurar_dia(dia) for i, dia in enumerate(DIAS_SEMANA)}
    salvar_horarios(grade)
    print("\n✔️ Memória de horários salva com sucesso!")
    return grade


def editar_horarios(grade_atual):
    escolha = input("[1] Refazer um Dia | [2] Refazer semana | [V] Voltar\n> ").strip().upper()
    if escolha == "2":
        return configurar_horarios()
    if escolha != "1":
        return grade_atual
    dia_idx = input("Dia [0-4]: ").strip()
    if dia_idx not in {"0", "1", "2", "3", "4"}:
        print("❌ Dia inválido.")
        return grade_atual
    grade_atual[dia_idx] = configurar_dia(DIAS_SEMANA[int(dia_idx)])
    salvar_horarios(grade_atual)
    print("✔️ Horário atualizado.")
    return grade_atual


def obter_turma_atual(grade):
    agora = datetime.now()
    dia_idx = str(agora.weekday())
    if dia_idx not in grade:
        return None
    minutos_agora = agora.hour * 60 + agora.minute
    for horario in HORARIOS_PADRAO:
        inicio, fim = horario.split(" - ")
        m_inicio = int(inicio[:2]) * 60 + int(inicio[3:])
        m_fim = int(fim[:2]) * 60 + int(fim[3:])
        if m_inicio - 10 <= minutos_agora <= m_fim:
            turma = grade[dia_idx].get(horario, "AULA ATIVIDADE")
            return turma if turma != "AULA ATIVIDADE" else None
    return None


def pasta_excel_destino():
    custom = os.getenv("MEGATRON_DIR")
    if custom:
        path = Path(custom)
    elif os.name == "nt":
        path = Path.home() / "OneDrive" / "Documentos" / "ROBO SIEPE"
    else:
        path = Path.cwd() / "ROBO_SIEPE"
    path.mkdir(parents=True, exist_ok=True)
    return path


def _tentar_fechar_excel(nome_arquivo):
    if os.name != "nt" or win32com is None:
        return False
    try:
        excel = win32com.client.GetActiveObject("Excel.Application")
        nome_base = os.path.basename(nome_arquivo)
        for wb_aberto in excel.Workbooks:
            if wb_aberto.Name == nome_base:
                wb_aberto.Save()
                wb_aberto.Close()
                return True
    except Exception:
        return False
    return False


def salvar_no_excel(turma, faltosos, notas_alunos, dia_letivo=True):
    nome_arquivo = pasta_excel_destino() / "Faltas_Megatron.xlsx"
    agora = datetime.now()
    data_hoje, horario_agora = agora.strftime("%d/%m/%Y"), agora.strftime("%H:%M:%S")

    nomes_faltosos = "NÃO LETIVO" if not dia_letivo else (", ".join(faltosos) if faltosos else "Todos presentes")
    total_faltas = 0 if not dia_letivo else len(faltosos)
    obs_geral = " | ".join([f"{a}: {n['texto']}" for a, n in notas_alunos.items()])

    linha = pd.DataFrame([{
        "Data": data_hoje,
        "Horário": horario_agora,
        "Turma": turma,
        "Alunos Faltosos": nomes_faltosos,
        "Total Faltas": total_faltas,
        "Observações do Dia": obs_geral,
    }])

    for tentativa in range(3):
        try:
            if nome_arquivo.exists():
                wb_tmp = openpyxl.load_workbook(nome_arquivo)
                startrow = wb_tmp[turma].max_row if turma in wb_tmp.sheetnames else 0
                wb_tmp.close()
                with pd.ExcelWriter(nome_arquivo, mode="a", engine="openpyxl", if_sheet_exists="overlay") as writer:
                    linha.to_excel(writer, sheet_name=turma, index=False, header=(startrow == 0), startrow=startrow)
            else:
                linha.to_excel(nome_arquivo, sheet_name=turma, index=False)

            wb = openpyxl.load_workbook(nome_arquivo)
            aba_ind = f"{turma} - Diário"
            ws_ind = wb[aba_ind] if aba_ind in wb.sheetnames else wb.create_sheet(aba_ind)

            if ws_ind.max_row == 1 and ws_ind.cell(1, 1).value is None:
                ws_ind.cell(1, 1, "Nome do Aluno")
            if ws_ind.cell(1, 1).value != "Nome do Aluno":
                ws_ind.cell(1, 1, "Nome do Aluno")

            cadastrados = {ws_ind.cell(row=i, column=1).value for i in range(2, ws_ind.max_row + 1)}
            for aluno in banco_alunos.get(turma, []):
                if aluno not in cadastrados:
                    ws_ind.cell(row=ws_ind.max_row + 1, column=1, value=aluno)

            col_data = ws_ind.max_column + 1
            ws_ind.cell(1, col_data, data_hoje)

            faltosos_set = set(faltosos)
            for row in range(2, ws_ind.max_row + 1):
                nome = ws_ind.cell(row=row, column=1).value
                if not nome:
                    continue
                status = "F" if nome in faltosos_set else "P"
                celula = ws_ind.cell(row=row, column=col_data, value=status)
                if status == "F":
                    celula.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
                if nome in notas_alunos:
                    celula.comment = Comment(notas_alunos[nome]["texto"], "Megatron")
                    celula.fill = PatternFill(start_color="9BC2E6", end_color="9BC2E6", fill_type="solid")

            wb.save(nome_arquivo)
            wb.close()
            print(f"☁️ Planilha atualizada: {nome_arquivo}")
            if os.name == "nt":
                try:
                    os.startfile(nome_arquivo)  # type: ignore[attr-defined]
                except OSError:
                    pass
            return

        except PermissionError:
            print("⚠️ Arquivo Excel em uso. Tentando fechar automaticamente...")
            if not _tentar_fechar_excel(str(nome_arquivo)):
                print("⚠️ Não foi possível fechar automaticamente. Feche manualmente e tente novamente.")
                return
            time.sleep(1.5)
        except Exception as exc:
            print(f"❌ Erro ao salvar Excel (tentativa {tentativa + 1}/3): {exc}")
            if tentativa == 2:
                return


def adicionar_anotacao_avulsa(turma, nome_aluno, texto_nota):
    nome_arquivo = pasta_excel_destino() / "Faltas_Megatron.xlsx"
    data_hoje = datetime.now().strftime("%d/%m/%Y")
    if not nome_arquivo.exists():
        print("❌ Arquivo Excel ainda não existe. Faça um lançamento primeiro.")
        return

    wb = openpyxl.load_workbook(nome_arquivo)
    aba_ind = f"{turma} - Diário"
    if aba_ind not in wb.sheetnames:
        print(f"⚠️ A aba '{aba_ind}' ainda não existia. O Megatron está criando agora...")
        ws = wb.create_sheet(aba_ind)
        ws.cell(1, 1, "Nome do Aluno")
        for idx, aluno in enumerate(banco_alunos.get(turma, []), start=2):
            ws.cell(idx, 1, aluno)
    else:
        ws = wb[aba_ind]
    col_data = None
    for col in range(2, ws.max_column + 1):
        if ws.cell(1, col).value == data_hoje:
            col_data = col
            break
    if col_data is None:
        col_data = ws.max_column + 1
        ws.cell(1, col_data, data_hoje)

    for row in range(2, ws.max_row + 1):
        if ws.cell(row, 1).value == nome_aluno:
            cel = ws.cell(row, col_data)
            cel.comment = Comment(texto_nota, "Megatron - Resumo")
            cel.fill = PatternFill(start_color="9BC2E6", end_color="9BC2E6", fill_type="solid")
            wb.save(nome_arquivo)
            wb.close()
            print("✔️ Anotação gravada no Excel.")
            return

    wb.close()
    print("❌ Aluno não encontrado na planilha.")


def carregar_credenciais():
    if CRED_FILE.exists():
        try:
            dados = json.loads(CRED_FILE.read_text(encoding="utf-8"))
            return dados.get("cpf", ""), dados.get("senha", "")
        except Exception:
            pass
    return "", ""


def salvar_credenciais(cpf, senha):
    CRED_FILE.write_text(json.dumps({"cpf": cpf, "senha": senha}, ensure_ascii=False), encoding="utf-8")
    try:
        os.chmod(CRED_FILE, 0o600)
    except OSError:
        pass


def executar_siepe(meu_cpf, minha_senha, texto_busca_siepe, modo_lancamento, lista_digitada):
    if webdriver is None:
        raise RuntimeError("Selenium não está instalado neste ambiente.")

    driver = None
    dia_letivo = False
    try:
        driver = webdriver.Chrome()
        driver.maximize_window()
        driver.get("https://siepe.educacao.pe.gov.br/")
        radar = WebDriverWait(driver, 20)

        radar.until(EC.element_to_be_clickable((By.ID, "login"))).send_keys(meu_cpf)
        driver.find_element(By.ID, "senha").send_keys(minha_senha + Keys.RETURN)
        radar.until(EC.element_to_be_clickable((By.ID, "mnSupEdu"))).click()
        radar.until(EC.element_to_be_clickable((By.CLASS_NAME, "ac-icon-diario-classe"))).click()
        driver.execute_script("arguments[0].click();", radar.until(EC.element_to_be_clickable((By.XPATH, "//span[contains(text(), 'Pesquisar')]"))))

        linha_turma = radar.until(EC.presence_of_element_located((By.XPATH, f"//td[contains(text(), '{texto_busca_siepe}')]/parent::tr")))
        driver.execute_script("arguments[0].scrollIntoView({block: 'center'});", linha_turma)
        driver.execute_script("arguments[0].click();", linha_turma)

        sanfona_freq = None
        for acc in driver.find_elements(By.CLASS_NAME, "DiarioClasseAccordion"):
            if "divFrequenciaDiaria" in (acc.get_attribute("onclick") or ""):
                sanfona_freq = acc
                break
        if sanfona_freq is None:
            raise RuntimeError("Painel de frequência diária não encontrado no SIEPE.")

        driver.execute_script("arguments[0].click();", sanfona_freq)
        caixas_existem = driver.find_elements(By.XPATH, "//input[contains(@id, 'FNJ')]")
        dia_letivo = len(caixas_existem) > 0

        if not dia_letivo:
            return dia_letivo

        if modo_lancamento == "P":
            for chk in caixas_existem:
                if not chk.is_selected():
                    driver.execute_script("arguments[0].click();", chk)

        if lista_digitada:
            links = driver.find_elements(By.XPATH, "//a[contains(@onclick, 'mostrarInformaçõesAluno')]")
            mapa = {
                normalizar_texto(link.text): (link.get_attribute("onclick") or "").split("'")[1]
                for link in links
                if link.get_attribute("onclick") and "'" in link.get_attribute("onclick")
            }
            for nome in lista_digitada:
                key = normalizar_texto(nome)
                if key in mapa:
                    chks = driver.find_elements(By.XPATH, f"//input[@value='{mapa[key]}' and contains(@id, 'FNJ')]")
                    for chk in chks:
                        if (modo_lancamento == "P" and chk.is_selected()) or (modo_lancamento == "F" and not chk.is_selected()):
                            driver.execute_script("arguments[0].click();", chk)

        try:
            caixa_conclusao = driver.find_element(By.ID, "chkConclusaoFrequenciaDiaria")
            if not caixa_conclusao.is_selected():
                driver.execute_script("arguments[0].click();", caixa_conclusao)
        except Exception:
            pass

        btn = radar.until(EC.presence_of_element_located((By.ID, "btnGravarFrequenciaDisciplina")))
        driver.execute_script("arguments[0].click();", btn)
        print("🚀 Frequência gravada no SIEPE!")
        return dia_letivo

    except KeyboardInterrupt as exc:
        raise RuntimeError("Automação interrompida pelo usuário (Ctrl+C).") from exc
    except (TimeoutException, WebDriverException) as exc:
        raise RuntimeError(f"Falha de automação no SIEPE: {exc}") from exc
    finally:
        if driver is not None:
            driver.quit()


def selecionar_turma(grade_semanal, escolha):
    if escolha == "M":
        return prompt("👉 Qual turma? (TAB): ", completer=TurmaCompleter()).strip().upper()
    turma_auto = obter_turma_atual(grade_semanal)
    if not turma_auto:
        print("⏳ Nenhum horário letivo ativo agora. Use o modo Manual [M].")
        return None
    print(f"⏱️ Turma detectada pelo horário: {turma_auto}")
    return next((t for t in lista_turmas if t.startswith(turma_auto)), None)


def main():
    print("\n" + "🤖 MEGATRON ONLINE v6.0 - Edição Blindada ".center(60, "="))
    grade_semanal = carregar_horarios() or configurar_horarios()

    while True:
        exibir_tabela_horarios(grade_semanal)
        escolha = input("[Enter] Auto | [M] Manual | [A] Anotar | [E] Editar | [S] Sair\n> ").strip().upper()

        if escolha == "S":
            break
        if escolha == "E":
            grade_semanal = editar_horarios(grade_semanal)
            continue

        if escolha == "A":
            turma = prompt("👉 Qual turma para anotação? (TAB): ", completer=TurmaCompleter()).strip().upper()
            if " - " not in turma:
                print("❌ Selecione a turma no formato correto.")
                continue
            sigla = turma.split(" - ", 1)[0].strip()
            if sigla not in banco_alunos:
                print("❌ Turma sem alunos cadastrados.")
                continue
            nome = prompt("Nome do aluno (TAB): ", completer=AlunoCompleter(banco_alunos[sigla])).strip().upper()
            if nome not in banco_alunos[sigla]:
                print("⚠️ Aluno não encontrado.")
                continue
            nota = input(f"Escreva a nota para {nome}:\n> ").strip()
            if nota:
                adicionar_anotacao_avulsa(sigla, nome, nota)
            continue

        turma_escolhida = selecionar_turma(grade_semanal, escolha)
        if not turma_escolhida or " - " not in turma_escolhida:
            print("❌ Turma inválida.")
            continue

        sigla_turma = turma_escolhida.split(" - ", 1)[0].strip()
        texto_busca_siepe = turma_escolhida.split(" - ", 1)[1].strip()
        if sigla_turma not in banco_alunos:
            print("❌ Turma sem alunos cadastrados.")
            continue

        alunos = banco_alunos[sigla_turma]
        modo = input("Modo: [F] Faltosos ou [P] Presentes? ").strip().upper()
        modo = modo if modo in {"F", "P"} else "F"

        entrada = prompt(
            "📝 Digite os PRESENTES (TAB): " if modo == "P" else "📝 Digite os FALTOSOS (TAB): ",
            completer=MegatronCompleter(alunos),
        )
        lista_digitada = [n.strip().upper() for n in entrada.split(",") if n.strip()]
        faltosos_excel = [a for a in alunos if a not in lista_digitada] if modo == "P" else lista_digitada

        notas_do_dia = {}
        while True:
            nome_obs = prompt("Aluno para observação (TAB/Enter para sair): ", completer=AlunoCompleter(alunos)).strip().upper()
            if not nome_obs:
                break
            if nome_obs in alunos:
                texto = input(f"Nota para {nome_obs}:\n> ").strip()
                if texto:
                    notas_do_dia[nome_obs] = {"texto": texto}
            else:
                print("⚠️ Aluno não encontrado.")

        cpf, senha = carregar_credenciais()
        if not cpf or not senha:
            cpf = input("CPF: ").strip()
            senha = getpass.getpass("Senha: ")
            salvar = input("Salvar credenciais localmente? [s/N]: ").strip().lower() == "s"
            if salvar:
                salvar_credenciais(cpf, senha)

        dia_letivo = True
        try:
            dia_letivo = executar_siepe(cpf, senha, texto_busca_siepe, modo, lista_digitada)
            if not dia_letivo:
                print("⏭️ Não é dia letivo no SIEPE.")
        except RuntimeError as exc:
            print(f"❌ {exc}")

        salvar_no_excel(sigla_turma, faltosos_excel, notas_do_dia, dia_letivo=dia_letivo)
        print("✨ Missão concluída.\n")
        time.sleep(1)


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        print("\n🛑 Execução interrompida pelo usuário. Encerrando com segurança...")
